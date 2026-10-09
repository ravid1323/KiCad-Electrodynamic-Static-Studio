import sys, os
import wx
import ctypes
import tempfile
from pathlib import Path
import json
import threading
import shutil
import subprocess
import traceback
import requests
import time
import signal
import re
import numpy as np
from kipy import KiCad
from kipy.proto.common.types import KiCadObjectType
from kipy.board_types import Track, ArcTrack, Via, Pad, PadType, Zone, PadStackShape
from kipy.util.board_layer import canonical_name, is_copper_layer
from ui.dashboard_server import OpenEMSServer, status_file_path
from analysis import plotter, post_processing
import pyvista as pv
import pptx
from pptx import Presentation
from pptx.util import Inches



def write_status_atomic(file_path, payload):
    """
    Writes status JSON atomically to prevent reader crashes (e.g. Dashboard)
    during active write cycles.
    """
    tmp_path = file_path + ".tmp"
    try:
        with open(tmp_path, "w", encoding="utf-8") as f:
            json.dump(payload, f)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp_path, file_path)
    except Exception as e:
        print(f"[!] Atomic JSON write failed: {e}")

def openems_locator(txt_log=None):
    try:
        import openEMS
    except:
        import pip
        openems_path = shutil.which("openEMS")
        if txt_log:
            txt_log.AppendText("OpenEMS Python module not found\n")
        openems_python_path = os.path.join(os.path.dirname(openems_path), "python")
        python_version = ".".join(sys.version.split(".")[:2])
        python_version_package_requirement = f'''cp{"".join(sys.version.split(".")[:2])}'''
        packages = list(filter(lambda pkg: python_version_package_requirement in pkg, os.listdir(openems_python_path)))[::-1]
        if len(packages):
            if txt_log:
                txt_log.AppendText("Found OpenEMS Python module packages found, installing...\n")
            for package in packages:
                process = subprocess.Popen([sys.executable, "-m", "pip", "install", os.path.join(openems_python_path, package)], shell=True, creationflags=subprocess.CREATE_NO_WINDOW)
                process.wait()
            process = subprocess.Popen([sys.executable, "-m", "pip", "install", "h5py"], shell=True, creationflags=subprocess.CREATE_NO_WINDOW)
            process.wait()
            if txt_log:
                txt_log.AppendText("Installation finished. PLEASE RESTART THE PLUGIN\n")
        elif txt_log:
            txt_log.AppendText(
                f"The installed version of OpenEMS does not support the installed Python version ({python_version})\n")
            txt_log.AppendText("consider using WSL\n")



class PortConfigDialog(wx.Dialog):
    def __init__(self, parent, net_names, get_pads_callback, copper_layers, existing_data=None):
        super().__init__(parent, title="Configure Port", size=(750, 450))
        self.net_names = [""] + net_names
        self.get_pads_callback = get_pads_callback
        self.copper_layers = copper_layers
        self.port_data = existing_data or {}

        main_sizer = wx.BoxSizer(wx.VERTICAL)

        # Upper Config Row
        hbox_top = wx.BoxSizer(wx.HORIZONTAL)
        hbox_top.Add(wx.StaticText(self, label="Type:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.cb_type = wx.ComboBox(self, choices=["Lumped", "Waveguide (Rectangular)", "Microstrip (MSL)", "Coaxial",
                                                  "Coplanar (CPW)"], style=wx.CB_READONLY)
        self.cb_type.SetStringSelection(self.port_data.get("type", "Lumped"))
        hbox_top.Add(self.cb_type, 1, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 10)

        hbox_top.Add(wx.StaticText(self, label="Z0 (Ω):"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.tc_impedance = wx.TextCtrl(self, value=str(self.port_data.get("impedance", "50.0")))
        hbox_top.Add(self.tc_impedance, 1, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 10)

        hbox_top.Add(wx.StaticText(self, label="Mode:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.combo_port_mode = wx.Choice(self,
                                         choices=["Single-Ended", "Direct Differential", "Mixed-Mode Differential"])
        self.combo_port_mode.SetStringSelection(self.port_data.get("mode", "Direct Differential"))
        hbox_top.Add(self.combo_port_mode, 2, wx.ALIGN_CENTER_VERTICAL, 0)
        main_sizer.Add(hbox_top, 0, wx.EXPAND | wx.ALL, 10)

        # Terminals Area
        hbox_terminals = wx.BoxSizer(wx.HORIZONTAL)

        # Positive Terminal
        sb_pos = wx.StaticBox(self, label="Positive Terminal (+)")
        sbs_pos = wx.StaticBoxSizer(sb_pos, wx.VERTICAL)
        self.cb_net_pos = wx.ComboBox(sb_pos, choices=self.net_names, style=wx.CB_READONLY)
        self.cb_pad_pos = wx.ComboBox(sb_pos, choices=[], style=wx.CB_READONLY)
        sbs_pos.Add(wx.StaticText(sb_pos, label="Select Net:"), 0, wx.TOP, 2)
        sbs_pos.Add(self.cb_net_pos, 0, wx.EXPAND | wx.BOTTOM, 3)
        sbs_pos.Add(wx.StaticText(sb_pos, label="Select Pad:"), 0, wx.TOP, 2)
        sbs_pos.Add(self.cb_pad_pos, 0, wx.EXPAND | wx.BOTTOM, 2)
        hbox_terminals.Add(sbs_pos, 1, wx.EXPAND | wx.RIGHT, 5)

        # Negative Terminal
        sb_neg = wx.StaticBox(self, label="Negative Terminal (-)")
        sbs_neg = wx.StaticBoxSizer(sb_neg, wx.VERTICAL)
        self.cb_net_neg = wx.ComboBox(sb_neg, choices=self.net_names, style=wx.CB_READONLY)
        self.cb_pad_neg = wx.ComboBox(sb_neg, choices=[], style=wx.CB_READONLY)
        sbs_neg.Add(wx.StaticText(sb_neg, label="Select Net:"), 0, wx.TOP, 2)
        sbs_neg.Add(self.cb_net_neg, 0, wx.EXPAND | wx.BOTTOM, 3)
        sbs_neg.Add(wx.StaticText(sb_neg, label="Select Pad:"), 0, wx.TOP, 2)
        sbs_neg.Add(self.cb_pad_neg, 0, wx.EXPAND | wx.BOTTOM, 2)
        hbox_terminals.Add(sbs_neg, 1, wx.EXPAND | wx.LEFT, 5)

        main_sizer.Add(hbox_terminals, 0, wx.EXPAND | wx.ALL, 10)

        # Layers and MSL Config
        hbox_layers = wx.BoxSizer(wx.HORIZONTAL)
        hbox_layers.Add(wx.StaticText(self, label="Signal Layer:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.cb_signal_layer = wx.ComboBox(self, choices=self.copper_layers, style=wx.CB_READONLY)
        self.cb_signal_layer.SetStringSelection(self.port_data.get("signal_layer", self.copper_layers[0]))
        hbox_layers.Add(self.cb_signal_layer, 1, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 3)

        hbox_layers.Add(wx.StaticText(self, label="Ref Layer (GND):"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.cb_ref_layer = wx.ComboBox(self, choices=self.copper_layers, style=wx.CB_READONLY)
        default_ref = self.copper_layers[1] if len(self.copper_layers) > 1 else self.copper_layers[-1]
        self.cb_ref_layer.SetStringSelection(self.port_data.get("reference_layer", default_ref))
        hbox_layers.Add(self.cb_ref_layer, 1, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 3)
        main_sizer.Add(hbox_layers, 0, wx.EXPAND | wx.ALL, 10)

        hbox_msl = wx.BoxSizer(wx.HORIZONTAL)
        hbox_msl.Add(wx.StaticText(self, label="MSL Len (mm):"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 2)
        self.tc_msl_len = wx.TextCtrl(self, value=str(self.port_data.get("msl_length", "50.0")))
        hbox_msl.Add(self.tc_msl_len, 1, wx.EXPAND | wx.RIGHT, 5)
        hbox_msl.Add(wx.StaticText(self, label="Feed Shift:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 2)
        self.tc_feed = wx.TextCtrl(self, value=str(self.port_data.get("feed_shift", "4.48")))
        hbox_msl.Add(self.tc_feed, 1, wx.EXPAND | wx.RIGHT, 5)
        hbox_msl.Add(wx.StaticText(self, label="Meas Shift:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 2)
        self.tc_meas = wx.TextCtrl(self, value=str(self.port_data.get("meas_plane_shift", "16.67")))
        hbox_msl.Add(self.tc_meas, 1, wx.EXPAND)
        hbox_msl.Add(wx.StaticText(self, label="CPW Gap:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 2)
        self.tc_cpw_gap = wx.TextCtrl(self, value=str(self.port_data.get("cpw_gap", "0.25")))
        hbox_msl.Add(self.tc_cpw_gap, 1, wx.EXPAND)
        main_sizer.Add(hbox_msl, 0, wx.EXPAND | wx.ALL, 10)

        # OK / Cancel Buttons
        btn_sizer = self.CreateButtonSizer(wx.OK | wx.CANCEL)
        main_sizer.Add(btn_sizer, 0, wx.ALIGN_RIGHT | wx.ALL, 10)

        self.SetSizer(main_sizer)

        # Event Bindings
        self.cb_net_pos.Bind(wx.EVT_COMBOBOX, self.on_net_pos_change)
        self.cb_net_neg.Bind(wx.EVT_COMBOBOX, self.on_net_neg_change)
        self.combo_port_mode.Bind(wx.EVT_CHOICE, self.on_mode_change)

        # Pre-fill data if editing
        if existing_data:
            self.cb_net_pos.SetStringSelection(existing_data.get("positive_terminal", {}).get("net", ""))
            self.on_net_pos_change(None)
            self.cb_pad_pos.SetStringSelection(existing_data.get("positive_terminal", {}).get("pad", ""))

            if existing_data.get("mode") != "Single-Ended":
                self.cb_net_neg.SetStringSelection(existing_data.get("negative_terminal", {}).get("net", ""))
                self.on_net_neg_change(None)
                self.cb_pad_neg.SetStringSelection(existing_data.get("negative_terminal", {}).get("pad", ""))

        self.on_mode_change(None)

    def on_net_pos_change(self, event):
        self.cb_pad_pos.Clear()
        for display_name, kiid_val in self.get_pads_callback(self.cb_net_pos.GetValue()):
            self.cb_pad_pos.Append(display_name, kiid_val)
        if self.cb_pad_pos.GetCount() > 0: self.cb_pad_pos.SetSelection(0)

    def on_net_neg_change(self, event):
        self.cb_pad_neg.Clear()
        for display_name, kiid_val in self.get_pads_callback(self.cb_net_neg.GetValue()):
            self.cb_pad_neg.Append(display_name, kiid_val)
        if self.cb_pad_neg.GetCount() > 0: self.cb_pad_neg.SetSelection(0)

    def on_mode_change(self, event):
        selected_mode = self.combo_port_mode.GetStringSelection()
        is_diff = ("Single-Ended" not in selected_mode)
        self.cb_net_neg.Enable(is_diff)
        self.cb_pad_neg.Enable(is_diff)
        needs_ref = ("Direct Differential" not in selected_mode)
        self.cb_ref_layer.Enable(needs_ref)

    def get_port_data(self):
        def safe_float(val, default):
            try:
                return float(val)
            except ValueError:
                return default

        data = {
            "type": self.cb_type.GetValue(),
            "impedance": safe_float(self.tc_impedance.GetValue(), 50.0),
            "msl_length": safe_float(self.tc_msl_len.GetValue(), 50.0),
            "feed_shift": safe_float(self.tc_feed.GetValue(), 4.48),
            "meas_plane_shift": safe_float(self.tc_meas.GetValue(), 16.67),
            "cpw_gap": safe_float(self.tc_cpw_gap.GetValue(), 0.25),
            "mode": self.combo_port_mode.GetStringSelection(),
            "signal_layer": self.cb_signal_layer.GetValue(),
            "reference_layer": self.cb_ref_layer.GetValue(),
            "positive_terminal": {
                "net": self.cb_net_pos.GetValue(),
                "pad": self.cb_pad_pos.GetStringSelection(),
                "kiid": self.cb_pad_pos.GetClientData(
                    self.cb_pad_pos.GetSelection()) if self.cb_pad_pos.GetSelection() != wx.NOT_FOUND else ""
            }
        }
        if data["mode"] != "Single-Ended":
            data["negative_terminal"] = {
                "net": self.cb_net_neg.GetValue(),
                "pad": self.cb_pad_neg.GetStringSelection(),
                "kiid": self.cb_pad_neg.GetClientData(
                    self.cb_pad_neg.GetSelection()) if self.cb_pad_neg.GetSelection() != wx.NOT_FOUND else ""
            }
        return data

class DCNetworkConfigDialog(wx.Dialog):
    def __init__(self, parent, net_names, get_pads_callback, existing_data=None):
        super().__init__(parent, title="Configure DC Network", size=(650, 400))
        self.net_names = net_names
        self.get_pads_callback = get_pads_callback
        self.network_data = existing_data or {}

        main_sizer = wx.BoxSizer(wx.VERTICAL)

        # Network Name
        hbox_name = wx.BoxSizer(wx.HORIZONTAL)
        hbox_name.Add(wx.StaticText(self, label="DC Network Name (e.g., P5V_Core, TX_Line_1):"), 0,
                      wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.tc_network_name = wx.TextCtrl(self, value=self.network_data.get("network_name", ""))
        hbox_name.Add(self.tc_network_name, 1, wx.EXPAND)
        main_sizer.Add(hbox_name, 0, wx.EXPAND | wx.ALL, 10)

        # Voltage Source
        sb_source = wx.StaticBox(self, label="Voltage Source (Dirichlet - Fixed Voltage)")
        sbs_source = wx.StaticBoxSizer(sb_source, wx.HORIZONTAL)
        self.cb_vrm_net = wx.ComboBox(sb_source, choices=self.net_names, style=wx.CB_READONLY)
        self.cb_vrm_pad = wx.ComboBox(sb_source, choices=[], style=wx.CB_READONLY)
        self.tc_vrm_volt = wx.TextCtrl(sb_source,
                                       value=str(self.network_data.get("vrm_source", {}).get("voltage", "5.0")))
        sbs_source.Add(wx.StaticText(sb_source, label="Net:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        sbs_source.Add(self.cb_vrm_net, 1, wx.EXPAND | wx.ALL, 2)
        sbs_source.Add(wx.StaticText(sb_source, label="Pad:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        sbs_source.Add(self.cb_vrm_pad, 1, wx.EXPAND | wx.ALL, 2)
        sbs_source.Add(wx.StaticText(sb_source, label="Volts:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        sbs_source.Add(self.tc_vrm_volt, 0, wx.EXPAND | wx.ALL, 2)
        main_sizer.Add(sbs_source, 0, wx.EXPAND | wx.ALL, 5)

        # Return Path
        sb_return = wx.StaticBox(self, label="Return Sink (0V Reference)")
        sbs_return = wx.StaticBoxSizer(sb_return, wx.HORIZONTAL)
        self.cb_gnd_net = wx.ComboBox(sb_return, choices=self.net_names, style=wx.CB_READONLY)
        self.cb_gnd_pad = wx.ComboBox(sb_return, choices=[], style=wx.CB_READONLY)
        sbs_return.Add(wx.StaticText(sb_return, label="Net:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        sbs_return.Add(self.cb_gnd_net, 1, wx.EXPAND | wx.ALL, 2)
        sbs_return.Add(wx.StaticText(sb_return, label="Pad:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        sbs_return.Add(self.cb_gnd_pad, 1, wx.EXPAND | wx.ALL, 2)
        main_sizer.Add(sbs_return, 0, wx.EXPAND | wx.ALL, 5)

        # Active Load
        sb_load = wx.StaticBox(self, label="Active Load (Current Sink)")
        sbs_load = wx.StaticBoxSizer(sb_load, wx.VERTICAL)

        hbox_tail = wx.BoxSizer(wx.HORIZONTAL)
        self.cb_load_tail_net = wx.ComboBox(sb_load, choices=self.net_names, style=wx.CB_READONLY)
        self.cb_load_tail_pad = wx.ComboBox(sb_load, choices=[], style=wx.CB_READONLY)
        hbox_tail.Add(wx.StaticText(sb_load, label="Tail Net:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        hbox_tail.Add(self.cb_load_tail_net, 1, wx.EXPAND | wx.ALL, 2)
        hbox_tail.Add(wx.StaticText(sb_load, label="Tail Pad:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        hbox_tail.Add(self.cb_load_tail_pad, 1, wx.EXPAND | wx.ALL, 2)
        sbs_load.Add(hbox_tail, 0, wx.EXPAND)

        hbox_head = wx.BoxSizer(wx.HORIZONTAL)
        self.cb_load_head_net = wx.ComboBox(sb_load, choices=self.net_names, style=wx.CB_READONLY)
        self.cb_load_head_pad = wx.ComboBox(sb_load, choices=[], style=wx.CB_READONLY)
        self.tc_load_amp = wx.TextCtrl(sb_load,
                                       value=str(self.network_data.get("load_sink", {}).get("current_A", "1.0")))
        hbox_head.Add(wx.StaticText(sb_load, label="Head Net:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        hbox_head.Add(self.cb_load_head_net, 1, wx.EXPAND | wx.ALL, 2)
        hbox_head.Add(wx.StaticText(sb_load, label="Head Pad:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        hbox_head.Add(self.cb_load_head_pad, 1, wx.EXPAND | wx.ALL, 2)
        hbox_head.Add(wx.StaticText(sb_load, label="Amps:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        hbox_head.Add(self.tc_load_amp, 0, wx.EXPAND | wx.ALL, 2)
        sbs_load.Add(hbox_head, 0, wx.EXPAND)

        main_sizer.Add(sbs_load, 0, wx.EXPAND | wx.ALL, 5)

        btn_sizer = self.CreateButtonSizer(wx.OK | wx.CANCEL)
        main_sizer.Add(btn_sizer, 0, wx.ALIGN_RIGHT | wx.ALL, 10)
        self.SetSizer(main_sizer)

        # Bindings
        self._bind_net_pad(self.cb_vrm_net, self.cb_vrm_pad)
        self._bind_net_pad(self.cb_gnd_net, self.cb_gnd_pad)
        self._bind_net_pad(self.cb_load_tail_net, self.cb_load_tail_pad)
        self._bind_net_pad(self.cb_load_head_net, self.cb_load_head_pad)

        if existing_data:
            self._prefill(existing_data)

    def _prefill(self, data):
        def set_net_pad(net_cb, pad_cb, net_name, pad_kiid):
            if net_name:
                net_cb.SetStringSelection(net_name)
                # Manually trigger the pad update
                pad_cb.Clear()
                for disp_name, kiid in self.get_pads_callback(net_name):
                    pad_cb.Append(disp_name, kiid)
                # Select the correct pad by its kiid
                for i in range(pad_cb.GetCount()):
                    if pad_cb.GetClientData(i) == pad_kiid:
                        pad_cb.SetSelection(i)
                        break

        vrm = data.get("vrm_source", {})
        set_net_pad(self.cb_vrm_net, self.cb_vrm_pad, vrm.get("net"), vrm.get("pad_kiid"))

        gnd = data.get("gnd_sink", {})
        set_net_pad(self.cb_gnd_net, self.cb_gnd_pad, gnd.get("net"), gnd.get("pad_kiid"))

        load = data.get("load_sink", {})
        set_net_pad(self.cb_load_tail_net, self.cb_load_tail_pad, load.get("tail_net"), load.get("tail_pad_kiid"))
        set_net_pad(self.cb_load_head_net, self.cb_load_head_pad, load.get("head_net"), load.get("head_pad_kiid"))

    def _bind_net_pad(self, cb_net, cb_pad):
        def on_net_change(event):
            cb_pad.Clear()
            net_name = event.GetString()
            if net_name:
                for disp_name, kiid in self.get_pads_callback(net_name):
                    cb_pad.Append(disp_name, kiid)
                if cb_pad.GetCount() > 0: cb_pad.SetSelection(0)

        cb_net.Bind(wx.EVT_COMBOBOX, on_net_change)

    def get_network_data(self):
        def get_kiid(cb):
            sel = cb.GetSelection()
            return cb.GetClientData(sel) if sel != wx.NOT_FOUND else ""

        return {
            "network_name": self.tc_network_name.GetValue().strip() or "Unnamed_Network",
            "vrm_source": {
                "net": self.cb_vrm_net.GetValue(),
                "pad_kiid": get_kiid(self.cb_vrm_pad),
                "voltage": float(self.tc_vrm_volt.GetValue() or 5.0)
            },
            "gnd_sink": {
                "net": self.cb_gnd_net.GetValue(),
                "pad_kiid": get_kiid(self.cb_gnd_pad),
                "voltage": 0.0
            },
            "load_sink": {
                "tail_net": self.cb_load_tail_net.GetValue(),
                "tail_pad_kiid": get_kiid(self.cb_load_tail_pad),
                "head_net": self.cb_load_head_net.GetValue(),
                "head_pad_kiid": get_kiid(self.cb_load_head_pad),
                "current_A": float(self.tc_load_amp.GetValue() or 1.0)
            }
        }

class ParasiticsDialog(wx.Dialog):
    def __init__(self, parent, comp_data):
        super().__init__(parent, title=f"Edit Parasitics - {comp_data['reference']}", size=(320, 200))
        self.ref = comp_data['reference']

        main_sizer = wx.BoxSizer(wx.VERTICAL)
        grid = wx.FlexGridSizer(3, 2, 10, 10)

        self.tc_R = wx.TextCtrl(self, value=str(comp_data.get('parasitic_R', '')))
        self.tc_L = wx.TextCtrl(self, value=str(comp_data.get('parasitic_L', '')))
        self.tc_C = wx.TextCtrl(self, value=str(comp_data.get('parasitic_C', '')))

        # Disable the primary value field based on component type
        if self.ref.startswith('R'):
            self.tc_R.Disable()
        elif self.ref.startswith('L'):
            self.tc_L.Disable()
        elif self.ref.startswith('C'):
            self.tc_C.Disable()

        grid.Add(wx.StaticText(self, label="Parasitic R (e.g. 0.5R):"), 0, wx.ALIGN_CENTER_VERTICAL)
        grid.Add(self.tc_R, 1, wx.EXPAND)
        grid.Add(wx.StaticText(self, label="Parasitic L (e.g. 1.5nH):"), 0, wx.ALIGN_CENTER_VERTICAL)
        grid.Add(self.tc_L, 1, wx.EXPAND)
        grid.Add(wx.StaticText(self, label="Parasitic C (e.g. 0.5pF):"), 0, wx.ALIGN_CENTER_VERTICAL)
        grid.Add(self.tc_C, 1, wx.EXPAND)

        main_sizer.Add(grid, 1, wx.EXPAND | wx.ALL, 15)

        btn_sizer = self.CreateButtonSizer(wx.OK | wx.CANCEL)
        main_sizer.Add(btn_sizer, 0, wx.ALIGN_RIGHT | wx.ALL, 10)

        self.SetSizer(main_sizer)

    def get_values(self):
        return {
            'parasitic_R': self.tc_R.GetValue().strip(),
            'parasitic_L': self.tc_L.GetValue().strip(),
            'parasitic_C': self.tc_C.GetValue().strip()
        }

class SpiceMappingDialog(wx.Dialog):
    def __init__(self, parent, num_ports, available_nets):
        super().__init__(parent, title="Map SPICE Ports to Nets", size=(450, 300))
        self.mappings = []
        main_sizer = wx.BoxSizer(wx.VERTICAL)

        main_sizer.Add(wx.StaticText(self, label="Map the S-Parameter ports to their DC Network roles:"), 0, wx.ALL, 10)

        grid = wx.FlexGridSizer(num_ports, 3, 5, 10)
        for i in range(num_ports):
            port_label = wx.StaticText(self, label=f"Port {i + 1}:")
            cb_net = wx.ComboBox(self, choices=available_nets, style=wx.CB_READONLY)
            cb_role = wx.ComboBox(self, choices=["None", "Source", "Load"], style=wx.CB_READONLY)
            cb_role.SetSelection(0)

            grid.Add(port_label, 0, wx.ALIGN_CENTER_VERTICAL)
            grid.Add(cb_net, 1, wx.EXPAND)
            grid.Add(cb_role, 1, wx.EXPAND)
            self.mappings.append((cb_net, cb_role))

        main_sizer.Add(grid, 1, wx.EXPAND | wx.ALL, 10)
        btn_sizer = self.CreateButtonSizer(wx.OK | wx.CANCEL)
        main_sizer.Add(btn_sizer, 0, wx.ALIGN_RIGHT | wx.ALL, 10)
        self.SetSizer(main_sizer)

    def get_mappings(self):
        return [{"net": net.GetValue(), "role": role.GetValue()} for net, role in self.mappings]

class SimProgressDialog(wx.Dialog):
    def __init__(self, parent, title):
        # Initialize as a compact window
        super().__init__(parent, title=title, size=(450, 180), style=wx.DEFAULT_DIALOG_STYLE | wx.RESIZE_BORDER)
        self.is_aborted = False

        main_sizer = wx.BoxSizer(wx.VERTICAL)

        # Status Text
        self.lbl_status = wx.StaticText(self, label="Initializing FDTD Engine...\n\n")
        main_sizer.Add(self.lbl_status, 0, wx.ALL | wx.EXPAND, 10)

        # Progress Bar
        self.gauge = wx.Gauge(self, range=100)
        main_sizer.Add(self.gauge, 0, wx.ALL | wx.EXPAND, 10)

        # Control Buttons
        btn_sizer = wx.BoxSizer(wx.HORIZONTAL)
        self.btn_more = wx.ToggleButton(self, label="More Info ▼")
        self.btn_abort = wx.Button(self, label="Abort Simulation")

        btn_sizer.Add(self.btn_more, 0, wx.ALL, 5)
        btn_sizer.AddStretchSpacer()
        btn_sizer.Add(self.btn_abort, 0, wx.ALL, 5)
        main_sizer.Add(btn_sizer, 0, wx.EXPAND | wx.LEFT | wx.RIGHT, 5)

        # Collapsible Log Console
        self.txt_log = wx.TextCtrl(self, style=wx.TE_MULTILINE | wx.TE_READONLY)
        self.txt_log.SetFont(wx.Font(9, wx.FONTFAMILY_TELETYPE, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_NORMAL))
        self.txt_log.Hide()  # Hidden by default
        main_sizer.Add(self.txt_log, 1, wx.EXPAND | wx.ALL, 10)

        self.SetSizer(main_sizer)

        # Event Bindings
        self.btn_more.Bind(wx.EVT_TOGGLEBUTTON, self.on_toggle_info)
        self.btn_abort.Bind(wx.EVT_BUTTON, self.on_abort)
        self.Bind(wx.EVT_CLOSE, self.on_abort)

    def on_toggle_info(self, event):
        """Expands the window to reveal the openEMS output text."""
        if self.btn_more.GetValue():
            self.btn_more.SetLabel("Less Info ▲")
            self.txt_log.Show()
            self.SetSize((750, 500))
        else:
            self.btn_more.SetLabel("More Info ▼")
            self.txt_log.Hide()
            self.SetSize((450, 180))
        self.Layout()

    def on_abort(self, event):
        self.is_aborted = True

    def Update(self, value, newmsg):
        self.gauge.SetValue(value)
        self.lbl_status.SetLabel(newmsg)
        return not self.is_aborted

class EMSimDialog(wx.Dialog):
    @property
    def txt_output(self):
        """Maps legacy self.txt_output references directly to the new NetManagerTab console."""
        return self.tab_nets.txt_output

    def __init__(self, parent, kicad_client, board):
        super(EMSimDialog, self).__init__(
            parent,
            title="KiCad Electrodynamic & Static Studio",
            size=(960, 960),
            style=wx.DEFAULT_DIALOG_STYLE | wx.RESIZE_BORDER | wx.MAXIMIZE_BOX | wx.MINIMIZE_BOX
        )
        icon_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "resources", "icon_24x24.png")
        if os.path.exists(icon_path):
            icon = wx.Icon(icon_path, wx.BITMAP_TYPE_PNG)
            self.SetIcon(icon)
        self.kicad = kicad_client
        self.board = board
        self.server_config_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "server_config.json")
        self.server_config_setting = {"port": 50234, "chat_id": "", "bot_token": ""}
        if os.path.exists(self.server_config_path):
            try:
                with open(self.server_config_path, 'r', encoding='utf-8') as f:
                    self.server_config_setting = json.load(f)
            except Exception as error:
                wx.MessageBox(f"Server config has syntax issue\nThe telegram bot is inactive\nError info:\n{error}", "Info", wx.OK | wx.ICON_INFORMATION)
        else:
            with open(self.server_config_path, 'w', encoding='utf-8') as f:
                json.dump(self.server_config_setting, f, indent=4)
        self.server = OpenEMSServer(port=self.server_config_setting.get("port", 50234))
        self.server.start()
        self.telegram_bot = TelegramSimBot(
            bot_token=self.server_config_setting.get("bot_token", ""),
            chat_id=self.server_config_setting.get("chat_id", ""),
            status_file_path=status_file_path
        )
        self.telegram_bot.start_listening()
        self.SetLayoutDirection(wx.Layout_LeftToRight)
        self.parent = parent
        self.layer_listbox = None
        self.stackup_data = []
        self.target_nets_data = []
        self.last_export_path = ""

        self.InitUI()

        self.re_excitation = re.compile(r"(?i)excitation signal length is:\s+(\d+)")
        self.re_energy = re.compile(r"(?i)timestep:\s+(\d+).*?energy:\s*~?([0-9\.eE+-]+).*?([\-\d\.]+)\s*dB")
        self.re_time = re.compile(r"\[@\s*(?:(\d+)h)?\s*(?:(\d+)m)?\s*(\d+)s\]")
        self.re_speed = re.compile(r"\(([\d\.eE+-]+)\s*s/TS\)")

    def InitUI(self):
        self.notebook = wx.Notebook(self)

        # Instantiate the new Net Manager directly
        self.tab_nets = NetManagerTab(self.notebook, self.board)
        self.notebook.AddPage(self.tab_nets, "Net Extraction")

        self.tab_stackup = wx.Panel(self.notebook)
        self.notebook.AddPage(self.tab_stackup, "Stackup & Sim Setup")

        # --- הוספת הטאבים החדשים ---
        self.tab_sim_settings = SimSettingsTab(self.notebook)
        self.notebook.AddPage(self.tab_sim_settings, "Sim & Excitation")

        self.tab_mesh_settings = MeshSettingsTab(self.notebook, main_app_ref=self)
        self.notebook.AddPage(self.tab_mesh_settings, "Mesh Configuration")

        self.tab_components = ComponentManagerTab(self.notebook, self.board)
        self.notebook.AddPage(self.tab_components, "RLC Components")
        # -----------------------------

        # --- Tab 2: Stackup & Layers ---
        vbox2 = wx.BoxSizer(wx.VERTICAL)
        self.txt_stackup = wx.TextCtrl(self.tab_stackup, style=wx.TE_MULTILINE | wx.TE_READONLY)
        self.txt_stackup.SetFont(wx.Font(10, wx.FONTFAMILY_TELETYPE, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_NORMAL))
        vbox2.Add(self.txt_stackup, 1, wx.EXPAND | wx.ALL, 5)

        # --- TAB 3: Dynamic Ports Configuration ---
        self.tab_ports = DynamicPortsTab(self.notebook, self)
        self.notebook.AddPage(self.tab_ports, "Ports Configuration")

        self.tab_dc = DCAnalysisTab(self.notebook, self)
        self.notebook.AddPage(self.tab_dc, "DC Analysis")
        self.tab_dc.update_nets()

        self.tab_post_processing = PostProcessingTab(self.notebook, self)
        self.notebook.AddPage(self.tab_post_processing, "Post-Processing")

        self.tab_report = ReportGeneratorTab(self.notebook, self)
        self.notebook.AddPage(self.tab_report, "Report Generator")

        self.tab_settings = SettingsTab(self.notebook)
        self.notebook.AddPage(self.tab_settings, "Settings & Dependencies")

        # Sim Type Selection
        hbox_sim_type = wx.BoxSizer(wx.HORIZONTAL)
        hbox_sim_type.Add(wx.StaticText(self.tab_stackup, label="Solver Model Type: "), 0,
                          wx.ALIGN_CENTER_VERTICAL | wx.ALL, 5)
        self.combo_sim_type = wx.ComboBox(self.tab_stackup,
                                          choices=["2.5D (Thin Sheet - Fast)", "3D (Volumetric - Accurate)"],
                                          style=wx.CB_READONLY)
        self.combo_sim_type.SetSelection(0)  # Default to 2.5D
        self.combo_sim_type.SetToolTip(
            "2.5D models copper as 0-thickness sheets (faster meshing). 3D uses true copper thickness.")
        hbox_sim_type.Add(self.combo_sim_type, 1, wx.EXPAND | wx.ALL, 5)
        vbox2.Add(hbox_sim_type, 0, wx.EXPAND | wx.ALL, 5)

        vbox2.Add(wx.StaticText(self.tab_stackup, label="Simulation Target Layers (Hold Ctrl/Shift or Ctrl+A):"), 0,
                  wx.ALL, 5)
        self.layer_listbox = wx.ListBox(self.tab_stackup, style=wx.LB_EXTENDED)
        self.layer_listbox.Bind(wx.EVT_KEY_DOWN, self.OnLayerListboxKeyDown)
        vbox2.Add(self.layer_listbox, 1, wx.EXPAND | wx.ALL, 5)

        btn_sizer2 = wx.BoxSizer(wx.HORIZONTAL)
        btn_sync = wx.Button(self.tab_stackup, label="Sync Stackup via S-Expression")
        btn_sync.Bind(wx.EVT_BUTTON, self.OnSyncStackup)
        btn_sizer2.Add(btn_sync, 0, wx.ALL, 5)

        vbox2.Add(btn_sizer2, 0, wx.ALIGN_CENTER)
        self.tab_stackup.SetSizer(vbox2)

        # --- Main Layout ---
        main_sizer = wx.BoxSizer(wx.VERTICAL)
        main_sizer.Add(self.notebook, 1, wx.EXPAND | wx.ALL, 5)

        btn_sizer = wx.BoxSizer(wx.HORIZONTAL)

        btn_refresh = wx.Button(self, label="Refresh KiCad Data")
        btn_refresh.Bind(wx.EVT_BUTTON, self.OnRefresh)
        btn_refresh.SetToolTip("Reload nets and board data from KiCad if you made changes.")
        btn_sizer.Add(btn_refresh, 0, wx.ALL, 10)

        btn_export = wx.Button(self, label="Export to JSON")
        btn_export.Bind(wx.EVT_BUTTON, self.OnExportJson)
        btn_export.SetBackgroundColour(wx.Colour(220, 240, 220))
        btn_sizer.Add(btn_export, 0, wx.TOP | wx.BOTTOM | wx.RIGHT, 10)

        # --- NEW BUTTON ON MAIN BAR ---
        btn_mesh = wx.Button(self, label="Calculate Mesh")
        btn_mesh.Bind(wx.EVT_BUTTON, self.tab_mesh_settings.on_calculate_mesh)
        btn_mesh.SetBackgroundColour(wx.Colour(230, 240, 255))
        btn_mesh.SetToolTip("Calculate the FDTD cell count without running the heavy simulation.")
        btn_sizer.Add(btn_mesh, 0, wx.TOP | wx.BOTTOM | wx.RIGHT, 10)

        # View Geometry Button
        btn_view = wx.Button(self, label="View 3D Geometry")
        btn_view.Bind(wx.EVT_BUTTON, self.OnViewGeometry)
        btn_view.SetBackgroundColour(wx.Colour(255, 250, 205))
        btn_view.SetToolTip("Open AppCSXCAD to visually inspect the model before running the math.")
        btn_sizer.Add(btn_view, 0, wx.TOP | wx.BOTTOM | wx.RIGHT, 10)

        # ADD PARAVIEW BUTTON HERE
        btn_paraview = wx.Button(self, label="Animate in ParaView")
        btn_paraview.Bind(wx.EVT_BUTTON, self.OnLaunchParaView)
        btn_paraview.SetBackgroundColour(wx.Colour(255, 200, 255))
        btn_paraview.SetToolTip("Open ParaView to animate the 3D E-field and H-field wave propagation.")
        btn_sizer.Add(btn_paraview, 0, wx.TOP | wx.BOTTOM | wx.RIGHT, 10)

        # Run Simulation Button
        btn_run = wx.Button(self, label="Run Simulation")
        btn_run.Bind(wx.EVT_BUTTON, self.OnRunOpenEMS)
        btn_run.SetBackgroundColour(wx.Colour(220, 220, 255))
        btn_run.SetToolTip("Execute openEMS FDTD solver and plot S-Parameters.")
        btn_sizer.Add(btn_run, 0, wx.TOP | wx.BOTTOM | wx.RIGHT, 10)

        btn_sizer.AddStretchSpacer(1)

        btn_close = wx.Button(self, label="Close")
        btn_close.Bind(wx.EVT_BUTTON, self.OnClose)
        btn_sizer.Add(btn_close, 0, wx.ALL, 10)

        main_sizer.Add(btn_sizer, 0, wx.EXPAND)
        self.SetSizer(main_sizer)

    def get_net_names(self):
        nets = self.board.get_nets()
        return sorted([net.name for net in nets if net.name])

    def get_layers(self):
        content = self.board.get_as_string()
        stackup_start = content.find('(stackup')
        if stackup_start == -1:
            self.txt_stackup.AppendText("No (stackup ...) block found in board data.\n")
            return

        depth = 0
        stackup_end = -1
        for i in range(stackup_start, len(content)):
            if content[i] == '(':
                depth += 1
            elif content[i] == ')':
                depth -= 1
                if depth == 0:
                    stackup_end = i
                    break

        stackup_content = content[stackup_start:stackup_end]
        layer_blocks = re.split(r'\((?:layer|sublayer)\s+"([^"]+)"', stackup_content)
        return layer_blocks

    def update_port_layer_choices(self):
        """Stores the dynamic copper layers from the stackup for the Port Configuration Dialog."""
        self.active_copper_layers = [l["layer_name"] for l in self.stackup_data if l.get("type") == "Copper"]

        if not self.active_copper_layers:
            self.active_copper_layers = ["F.Cu", "B.Cu"]  # Fallback


    def get_pads_for_net(self, net_name):
        if not net_name:
            return []

        unique_pads = set()

        for footprint in self.board.get_footprints():

            ref = "Unknown"
            if hasattr(footprint, 'reference_field') and footprint.reference_field:
                txt = footprint.reference_field.text
                ref = txt.value if hasattr(txt, 'value') else str(txt)
            elif hasattr(footprint, 'reference'):
                ref = str(footprint.reference)

            # safe pad access
            pads_list = getattr(footprint, 'pads', None)
            if pads_list is None and hasattr(footprint, 'definition'):
                pads_list = getattr(footprint.definition, 'pads', [])

            if not pads_list:
                continue

            for pad in pads_list:
                if hasattr(pad, 'net') and pad.net and pad.net.name == net_name:
                    pad_num = str(getattr(pad, 'number', ''))
                    pad_display_name = f"{ref}.{pad_num}" if ref != "Unknown" else pad_num
                    kiid_value = pad.id.value if hasattr(pad.id, 'value') else str(pad.id)
                    unique_pads.add((pad_display_name, kiid_value))

        return sorted(list(unique_pads), key=lambda x: x[0])

    def OnRefresh(self, event):
        try:
            self.board = self.kicad.get_board()

            # Sync child tabs safely
            self.tab_nets.board = self.board
            self.tab_nets.LoadNets()
            self.tab_components.board = self.board
            self.tab_components.LoadComponents()

            self.tab_dc.update_nets()

            self.txt_output.Clear()
            self.target_nets_data = []
            self.txt_stackup.Clear()
            self.txt_stackup.AppendText("Board data refreshed from KiCad.\nPlease 'Sync Stackup' again.\n")
            self.stackup_data = []
            if self.layer_listbox: self.layer_listbox.Clear()
        except Exception as e:
            wx.MessageBox(f"Failed to refresh data:\n{e}", "Error", wx.OK | wx.ICON_ERROR)

    def OnSyncStackup(self, event):
        self.txt_stackup.Clear()
        self.txt_stackup.AppendText("Fetching Stackup via S-Expression Parse...\n")
        self.txt_stackup.AppendText("-" * 80 + "\n")
        self.txt_stackup.AppendText(f"{'Layer':<15} | {'Type':<12} | {'Thick(mm)':<10} | {'Material (Er, TanD)'}\n")
        self.txt_stackup.AppendText("-" * 80 + "\n")

        copper_layers = []
        self.stackup_data = []

        try:
            content = self.board.get_as_string()
            stackup_start = content.find('(stackup')
            if stackup_start == -1:
                self.txt_stackup.AppendText("No (stackup ...) block found in board data.\n")
                return

            depth = 0
            stackup_end = -1
            for i in range(stackup_start, len(content)):
                if content[i] == '(':
                    depth += 1
                elif content[i] == ')':
                    depth -= 1
                    if depth == 0:
                        stackup_end = i
                        break

            stackup_content = content[stackup_start:stackup_end]
            layer_blocks = re.split(r'\((?:layer|sublayer)\s+"([^"]+)"', stackup_content)

            for i in range(1, len(layer_blocks), 2):
                name = layer_blocks[i]
                props = layer_blocks[i + 1]

                thick_m = re.search(r'\(thickness\s+([0-9.]+)', props)
                type_m = re.search(r'\(type\s+"?([^"\s\)]+)"?', props)
                mat_m = re.search(r'\(material\s+"([^"]+)"\)', props)
                er_m = re.search(r'\(epsilon_r\s+([0-9.]+)', props)
                loss_m = re.search(r'\(loss_tangent\s+([0-9.]+)', props)

                if thick_m:
                    thick = float(thick_m.group(1))
                else:
                    thick_matches = re.findall(r'\(thickness\s+([0-9.]+)', props)
                    thick = sum(float(t) for t in thick_matches) if thick_matches else 0.0

                l_type = type_m.group(1) if type_m else "Unknown"
                is_copper = (l_type.lower() == 'copper')
                clean_type = "Copper" if is_copper else "Dielectric"

                if is_copper:
                    copper_layers.append(name)

                mat_parts = []
                mat_name = mat_m.group(1) if mat_m else ("Copper" if is_copper else "")
                if mat_name: mat_parts.append(mat_name)

                er_val = float(er_m.group(1)) if er_m else (1.0 if not is_copper else 0.0)
                loss_val = float(loss_m.group(1)) if loss_m else 0.02

                if er_m: mat_parts.append(f"Er: {er_val}")
                if loss_m: mat_parts.append(f"TanD: {loss_val}")

                mat_str = " | ".join(mat_parts) if mat_parts else ""
                self.txt_stackup.AppendText(f"{name:<15} | {clean_type:<12} | {thick:<10.4f} | {mat_str}\n")

                self.stackup_data.append({
                    "layer_name": name,
                    "type": clean_type,
                    "thickness_mm": thick,
                    "material": mat_name,
                    "epsilon_r": er_val,
                    "loss_tangent": loss_val
                })

            if self.layer_listbox:
                self.layer_listbox.SetItems(copper_layers)
                for i in range(self.layer_listbox.GetCount()):
                    self.layer_listbox.SetSelection(i)

            self.update_port_layer_choices()

            self.txt_stackup.AppendText("\n--- Stackup extraction complete ---\n")

        except Exception as e:
            self.txt_stackup.AppendText(f"Error fetching stackup: {e}\n")
            self.txt_stackup.AppendText(traceback.format_exc())


    def OnLayerListboxKeyDown(self, event):
        if event.ControlDown() and event.GetKeyCode() == ord('A'):
            for i in range(self.layer_listbox.GetCount()):
                self.layer_listbox.SetSelection(i)
        else:
            event.Skip()

    def extract_selected_nets(self):
        selected_net_names = self.tab_nets.GetSelectedNets()
        self.target_nets_data = []
        self.tab_nets.txt_output.Clear()

        if not selected_net_names:
            return

        self.tab_nets.txt_output.AppendText(f"Scanning Nets: {', '.join(selected_net_names)}...\n\n")

        all_nets = self.board.get_nets()
        selected_net_objs = [n for n in all_nets if n.name in selected_net_names]

        if not selected_net_objs:
            self.tab_nets.txt_output.AppendText("No matching net objects found.\n")
            return

        # Hardcode generic clearance and ignore Kipy design rules
        clearance_mm = 0.25
        extracted_data = {name: {'tracks': [], 'vias': [], 'tht_pads': [], 'zones': [], 'smd_pads': []} for name in selected_net_names}

        # --- MAP ALL PADS TO THEIR FOOTPRINT LAYER & REFERENCE ---
        pad_to_fp_layer = {}
        pad_to_fp_ref = {}
        try:
            for fp in self.board.get_footprints():
                fp_layer = canonical_name(fp.layer)
                ref = "Unknown"
                if hasattr(fp, 'reference_field') and fp.reference_field:
                    txt = fp.reference_field.text
                    ref = txt.value if hasattr(txt, 'value') else str(txt)
                elif hasattr(fp, 'reference'):
                    ref = str(fp.reference)

                pads_list = getattr(fp, 'pads', None)
                if pads_list is None and hasattr(fp, 'definition'):
                    pads_list = getattr(fp.definition, 'pads', [])

                for p in pads_list:
                    p_kiid = p.id.value if hasattr(p.id, 'value') else str(p.id)
                    pad_to_fp_layer[p_kiid] = fp_layer
                    pad_to_fp_ref[p_kiid] = ref
        except Exception as e:
            self.tab_nets.txt_output.AppendText(f"[!] Warning: Could not map footprint layers ({e})\n")

        # --- EXTRACT GEOMETRY ---
        try:
            items = self.board.get_items_by_net(
                nets=selected_net_objs,
                types=[
                    KiCadObjectType.KOT_PCB_TRACE,
                    KiCadObjectType.KOT_PCB_ARC,
                    KiCadObjectType.KOT_PCB_VIA,
                    KiCadObjectType.KOT_PCB_PAD,
                    KiCadObjectType.KOT_PCB_ZONE
                ]
            )

            for item in items:
                item_net_name = getattr(getattr(item, 'net', None), 'name', None)
                if item_net_name not in extracted_data:
                    continue

                if isinstance(item, Track) or isinstance(item, ArcTrack):
                    x1 = round(item.start.x / 1_000_000.0, 4)
                    y1 = round(item.start.y / 1_000_000.0, 4)
                    x2 = round(item.end.x / 1_000_000.0, 4)
                    y2 = round(item.end.y / 1_000_000.0, 4)
                    width_mm = round(item.width / 1_000_000.0, 4)
                    layer_name = canonical_name(item.layer)
                    extracted_data[item_net_name]['tracks'].append({
                        "start_x": x1, "start_y": y1, "end_x": x2, "end_y": y2,
                        "width_mm": width_mm, "layer": layer_name
                    })
                elif isinstance(item, Via):
                    x = round(item.position.x / 1_000_000.0, 4)
                    y = round(item.position.y / 1_000_000.0, 4)
                    size = round(item.diameter / 1_000_000.0, 4)
                    try:
                        s_layer = canonical_name(item.padstack.drill.start_layer)
                        e_layer = canonical_name(item.padstack.drill.end_layer)
                    except:
                        s_layer, e_layer = "", ""

                    try:
                        drill_mm = item.drill_diameter / 1_000_000.0
                    except:
                        drill_mm = 0.0

                    extracted_data[item_net_name]['vias'].append({
                        "x": x, "y": y, "size_mm": size, "drill_mm": drill_mm,
                        "start_layer": s_layer, "end_layer": e_layer, "net_name": item_net_name
                    })

                elif isinstance(item, Pad):
                    if item.pad_type in [PadType.PT_PTH, PadType.PT_NPTH, PadType.PT_SMD]:
                        x = round(item.position.x / 1_000_000.0, 4)
                        y = round(item.position.y / 1_000_000.0, 4)
                        pad_kiid = item.id.value if hasattr(item.id, 'value') else str(item.id)
                        fp_ref = pad_to_fp_ref.get(pad_kiid, "")
                        base_layer = pad_to_fp_layer.get(pad_kiid, "F.Cu")
                        pad_num = str(getattr(item, 'number', ""))
                        pad_id = f"{fp_ref}.{pad_num}" if (fp_ref and fp_ref != "None") else pad_num

                        try:
                            drill_dia_nm = item.padstack.drill.diameter.x
                            drill_mm = (drill_dia_nm / 1_000_000.0) if drill_dia_nm > 0 else 0.0
                        except:
                            drill_mm = 0.0

                        size = 1.5
                        try:
                            size_x = item.padstack.copper_layers[0].size.x / 1_000_000.0
                            size_y = item.padstack.copper_layers[0].size.y / 1_000_000.0
                            size = max(size_x, size_y)
                        except:
                            pass

                        pad_shapes = []
                        try:
                            if hasattr(item, 'padstack') and item.padstack.copper_layers:
                                for cu_layer in item.padstack.copper_layers:
                                    if is_copper_layer(cu_layer.layer):
                                        l_name = canonical_name(cu_layer.layer)
                                        if item.pad_type == PadType.PT_SMD:
                                            l_name = base_layer
                                        shape_name = PadStackShape.Name(cu_layer.shape)
                                        offset_x, offset_y = 0.0, 0.0
                                        if hasattr(cu_layer, 'offset'):
                                            offset_x = round(cu_layer.offset.x / 1_000_000.0, 4)
                                            offset_y = round(cu_layer.offset.y / 1_000_000.0, 4)
                                        pad_shapes.append({
                                            "layer": l_name,
                                            "shape": shape_name,
                                            "size_x_mm": round(cu_layer.size.x / 1_000_000.0, 4),
                                            "size_y_mm": round(cu_layer.size.y / 1_000_000.0, 4),
                                            "offset_x": offset_x,
                                            "offset_y": offset_y
                                        })
                        except Exception:
                            pass

                        if not pad_shapes:
                            pad_shapes.append({
                                "layer": base_layer,
                                "shape": "PSS_RECT",
                                "size_x_mm": size,
                                "size_y_mm": size,
                                "offset_x": 0.0,
                                "offset_y": 0.0
                            })

                        pad_data = {
                            "kiid": pad_kiid, "id": pad_id, "pad_num": pad_num,
                            "parent_ref": fp_ref, "net_name": item_net_name,
                            "x": x, "y": y, "size_mm": size, "drill_mm": drill_mm,
                            "layer": base_layer,
                            "pad_shapes": pad_shapes
                        }

                        if item.pad_type == PadType.PT_SMD:
                            extracted_data[item_net_name]['smd_pads'].append(pad_data)
                        else:
                            extracted_data[item_net_name]['tht_pads'].append(pad_data)

                elif isinstance(item, Zone):
                    for layer_enum, polys in item.filled_polygons.items():
                        layer_name = canonical_name(layer_enum)
                        for poly in polys:
                            outline_pts = []
                            for n in poly.outline:
                                if n.has_point:
                                    outline_pts.append({"x": round(n.point.x / 1_000_000.0, 4),
                                                        "y": round(n.point.y / 1_000_000.0, 4)})
                                elif n.has_arc:
                                    outline_pts.append({"x": round(n.arc.start.x / 1_000_000.0, 4),
                                                        "y": round(n.arc.start.y / 1_000_000.0, 4)})
                                    outline_pts.append({"x": round(n.arc.mid.x / 1_000_000.0, 4),
                                                        "y": round(n.arc.mid.y / 1_000_000.0, 4)})
                                    outline_pts.append({"x": round(n.arc.end.x / 1_000_000.0, 4),
                                                        "y": round(n.arc.end.y / 1_000_000.0, 4)})

                            holes_data = []
                            for hole in poly.holes:
                                hole_pts = []
                                for n in hole:
                                    if n.has_point:
                                        hole_pts.append({"x": round(n.point.x / 1_000_000.0, 4),
                                                         "y": round(n.point.y / 1_000_000.0, 4)})
                                    elif n.has_arc:
                                        hole_pts.append({"x": round(n.arc.start.x / 1_000_000.0, 4),
                                                         "y": round(n.arc.start.y / 1_000_000.0, 4)})
                                        hole_pts.append({"x": round(n.arc.mid.x / 1_000_000.0, 4),
                                                         "y": round(n.arc.mid.y / 1_000_000.0, 4)})
                                        hole_pts.append({"x": round(n.arc.end.x / 1_000_000.0, 4),
                                                         "y": round(n.arc.end.y / 1_000_000.0, 4)})
                                if hole_pts:
                                    holes_data.append(hole_pts)

                            if outline_pts:
                                extracted_data[item_net_name]['zones'].append({
                                    "layer": layer_name, "points": outline_pts, "holes": holes_data
                                })

            for net_name in selected_net_names:
                data = extracted_data[net_name]
                self.target_nets_data.append({
                    "name": net_name,
                    "clearance_mm": clearance_mm,
                    "tracks": data['tracks'],
                    "vias": data['vias'],
                    "tht_pads": data['tht_pads'],
                    "zones": data['zones'],
                    "smd_pads": data['smd_pads']
                })

                self.tab_nets.txt_output.AppendText(
                    f"--- Extracted: {len(data['tracks'])} segments, {len(data['vias'])} vias, {len(data['zones'])} zones for {net_name} ---\n\n"
                )

        except Exception as e:
            err_msg = str(e).lower()
            if "busy" in err_msg:
                self.tab_nets.txt_output.AppendText("\n[!] KiCad API is busy.\n-> Please click on the PCB Editor window, press 'ESC' a few times to cancel any active tool, and try selecting the nets again.\n")
            else:
                self.tab_nets.txt_output.AppendText(f"\nError fetching items: {e}\n")
                self.tab_nets.txt_output.AppendText(traceback.format_exc())

    def OnExportJson(self, event):
        self._export_json(event=event)

    def _export_json(self, event=None):
        self.extract_selected_nets()
        def get_prop_dir(net_name, pad_kiid):
            px, py = None, None
            for net in self.target_nets_data:
                if net["name"] == net_name:
                    for p in net.get("smd_pads", []) + net.get("tht_pads", []):
                        if p["kiid"] == pad_kiid:
                            px, py = p["x"], p["y"]
                            break
                    if px is not None:
                        for t in net.get("tracks", []):
                            dx, dy = 0.0, 0.0
                            if abs(t["start_x"] - px) < 0.001 and abs(t["start_y"] - py) < 0.001:
                                dx = t["end_x"] - px
                                dy = t["end_y"] - py
                            elif abs(t["end_x"] - px) < 0.001 and abs(t["end_y"] - py) < 0.001:
                                dx = t["start_x"] - px
                                dy = t["start_y"] - py

                            if dx != 0.0 or dy != 0.0:
                                if abs(dx) > abs(dy):
                                    return ('x', 1 if dx > 0 else -1)
                                else:
                                    return ('y', 1 if dy > 0 else -1)
            return ('x', 1)

        if not self.stackup_data:
            wx.MessageBox("Please Sync Stackup first!", "Warning", wx.OK | wx.ICON_WARNING)
            return

        if not self.target_nets_data:
            wx.MessageBox("Please Select at least one Net first!", "Warning", wx.OK | wx.ICON_WARNING)
            return

        selected_layers = [self.layer_listbox.GetString(i) for i in self.layer_listbox.GetSelections()]

        current_global_clearance = 0.25
        if getattr(self, 'kicad_is_legacy', False) and hasattr(self, 'txt_antipad_clearance'):
            try:
                current_global_clearance = float(self.txt_antipad_clearance.GetValue())
            except ValueError:
                pass

        manual_ports = {}
        for idx, port_data in enumerate(self.tab_ports.configured_ports):
            port_name = f"port_{idx + 1}"

            net_pos = port_data["positive_terminal"]["net"]
            p_kiid = port_data["positive_terminal"]["kiid"]
            p_axis, p_sign = get_prop_dir(net_pos, p_kiid)
            port_data["positive_terminal"]["prop_dir"] = p_axis
            port_data["positive_terminal"]["direction_sign"] = p_sign

            if port_data["mode"] != "Single-Ended":
                net_neg = port_data["negative_terminal"]["net"]
                n_kiid = port_data["negative_terminal"]["kiid"]
                n_axis, n_sign = get_prop_dir(net_neg, n_kiid)
                port_data["negative_terminal"]["prop_dir"] = n_axis
                port_data["negative_terminal"]["direction_sign"] = n_sign

            manual_ports[port_name] = port_data

        is_2_5d = "2.5D" in self.combo_sim_type.GetValue()

        try:
            project_dir = os.path.join(self.kicad.get_project(self.board.document).path, "EM Simuation Studio")
        except Exception:
            project_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "EM Simuation Studio")

        output_dir = os.path.join(project_dir, f'''em_simulation_results_{time.strftime("%Y-%m-%d_%H-%M-%S", time.localtime(time.time()))}''')
        os.makedirs(output_dir, exist_ok=True)

        sim_data_dir = os.path.join(output_dir, "fdtd_results")
        dc_data_dir = os.path.join(output_dir, "dc_results")
        os.makedirs(sim_data_dir, exist_ok=True)
        os.makedirs(dc_data_dir, exist_ok=True)

        edge_cuts_data = []
        try:
            for shape in self.board.get_shapes():
                if hasattr(shape, 'layer') and canonical_name(shape.layer) == "Edge.Cuts":
                    sx, sy, ex, ey = None, None, None, None
                    if hasattr(shape, 'top_left') and hasattr(shape, 'bottom_right'):
                        sx, sy = shape.bottom_right.x, shape.bottom_right.y
                        ex, ey = shape.top_left.x, shape.top_left.y
                    if sx is not None and ex is not None:
                        edge_cuts_data.append({
                            "start_x": round(sx / 1_000_000.0, 4), "start_y": round(sy / 1_000_000.0, 4),
                            "end_x": round(ex / 1_000_000.0, 4), "end_y": round(ey / 1_000_000.0, 4)
                        })
        except Exception as e:
            self.txt_output.AppendText(f"[!] Warning: Could not extract Edge.Cuts ({e})\n")

        export_data = {
            "project_info": self.board.get_project().name,
            "simulation_dir": sim_data_dir,
            "dc_simulation_dir": dc_data_dir,
            "simulation_type": "2.5D" if is_2_5d else "3D",
            "default_antipad_clearance_mm": current_global_clearance,
            "edge_cuts": edge_cuts_data,
            "stackup": self.stackup_data,
            "target_nets": self.target_nets_data,
            "sim_layers": selected_layers,
            "manual_ports": manual_ports
        }

        export_data.update(self.tab_sim_settings.get_data())
        export_data.update(self.tab_mesh_settings.get_data())
        export_data.update(self.tab_components.get_data())
        export_data.update(self.tab_dc.get_data())

        json_filename = os.path.join(output_dir, "simulation_metadata.json")
        try:
            with open(json_filename, 'w', encoding='utf-8') as f:
                json.dump(export_data, f, indent=4)
            self.last_export_path = json_filename
            wx.MessageBox(f"Successfully exported JSON model to:\n{json_filename}", "Success",
                          wx.OK | wx.ICON_INFORMATION)
        except IOError as e:
            wx.MessageBox(f"Failed to save file:\n{e}", "Error", wx.OK | wx.ICON_ERROR)

    def OnViewGeometry(self, event):
        if not self.last_export_path or not os.path.exists(self.last_export_path):
            wx.MessageBox("Please 'Export to JSON' first.", "Info", wx.OK | wx.ICON_INFORMATION)
            return

        script_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "solvers", "run_openems.py")

        if not os.path.exists(script_path):
            wx.MessageBox(f"Cannot find simulation script at:\n{script_path}", "Error", wx.OK | wx.ICON_ERROR)
            return

        try:
            subprocess.Popen([sys.executable, script_path, self.last_export_path, "--view"], creationflags=subprocess.CREATE_NO_WINDOW)
        except Exception as e:
            wx.MessageBox(f"Failed to launch View:\n{e}", "Launch Error", wx.OK | wx.ICON_ERROR)

    def OnRunOpenEMS(self, event):
        if not self.last_export_path or not os.path.exists(self.last_export_path):
            wx.MessageBox("Please 'Export to JSON' first.", "Info", wx.OK | wx.ICON_INFORMATION)
            return

        script_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "solvers", "run_openems.py")
        cmd = [sys.executable, "-u", script_path, self.last_export_path, "--run"]

        try:
            num_physical_ports = sum(
                [2 if p.get("mode") != "Single-Ended" else 1 for p in self.tab_ports.configured_ports])
            assume_symmetry = self.tab_sim_settings.chk_symmetry.GetValue()
            if num_physical_ports > 2:
                self.total_runs = (num_physical_ports // 2) if assume_symmetry else num_physical_ports
            else:
                self.total_runs = 1
        except Exception:
            self.total_runs = 1

        self.current_run = 0
        self.run_durations = []
        self.current_run_start_time = None

        try:
            target_db = float(self.tab_sim_settings.cmb_energy_limit.GetValue())
        except ValueError:
            target_db = -30.0

        run_status_str = f" (Run 1/{self.total_runs})" if self.total_runs > 1 else ""
        reset_payload = {
            "status": "Starting",
            "text_status": f"Starting FDTD Engine... 🚀{run_status_str}",
            "bg_color": "#1e3a8a",
            "energy_db": 0.0,
            "target_db": target_db,
            "eta": "ETA: Initializing...",
            "progress_pct": 0,
            "plot_path": ""
        }

        try:
            write_status_atomic(status_file_path, reset_payload)
        except Exception:
            pass

        self.txt_output.AppendText("\n--- Starting Simulation ---\n")
        self.energy_history = []

        run_status_str = f" (Run 1/{getattr(self, 'total_runs', 1)})" if getattr(self, 'total_runs', 1) > 1 else ""
        self.progress_dialog = SimProgressDialog(self, "openEMS Simulation")

        self.progress_dialog.Show()

        self.sim_thread = SimulationThread(cmd, self.OnSimulationOutput, self.OnSimulationDone)
        self.sim_thread.start()

    def OnSimulationOutput(self, chunk_str):
        if getattr(self, 'progress_dialog', None):
            wx.CallAfter(self.progress_dialog.txt_log.AppendText, chunk_str)
        else:
            wx.CallAfter(self.txt_output.AppendText, chunk_str)

        if not hasattr(self, '_sim_progress'):
            self._sim_progress = 0
            self._sim_msg = "Initializing FDTD Engine...\nThis may take a minute."
            self._sim_metric = 0.0
            self._sim_target = -30.0
            self._sim_eta_str = "ETA: Calculating..."
            self._sim_run_str = ""

        latest_payload = None
        found_update = False

        for line_str in chunk_str.splitlines():
            exc_match = self.re_excitation.search(line_str)
            if exc_match:
                self.excitation_length = int(exc_match.group(1))

            if "[*] Executing FDTD Run" in line_str:
                if getattr(self, 'current_run', 0) > 0 and getattr(self, 'current_run_start_time', None):
                    self.run_durations.append(time.time() - self.current_run_start_time)

                self.current_run = getattr(self, 'current_run', 0) + 1
                self.current_run_start_time = time.time()
                self.energy_history = []
                self.excitation_length = 0
                self.auto_aborted = False

            if not hasattr(self, 'prev_linear_energy'):
                self.prev_linear_energy = 0.0

            match = self.re_energy.search(line_str)

            if match:
                found_update = True
                current_step = int(match.group(1))
                linear_energy = float(match.group(2))
                current_db = -abs(float(match.group(3)))

                try:
                    target_db = float(self.tab_sim_settings.cmb_energy_limit.GetValue())
                    max_ts = int(self.tab_sim_settings.txt_max_timesteps.GetValue())
                except ValueError:
                    target_db = -30.0
                    max_ts = 100000

                excite_val = self.tab_sim_settings.cmb_excite_type.GetValue()
                is_step = "Step" in excite_val or "Sinusoid" in excite_val

                self.linear_energy_history = getattr(self, 'linear_energy_history', [])
                self.linear_energy_history.append((current_step, linear_energy))
                self.linear_energy_history = self.linear_energy_history[-1000:]

                if is_step:
                    if len(self.linear_energy_history) == 1000:
                        e_curr = self.linear_energy_history[-1][1]
                        e_past = self.linear_energy_history[-1000][1]
                        fractional_change = abs(e_curr - e_past) / (e_curr + 1e-30)
                        active_metric = float(10.0 * np.log10(fractional_change + 1e-30))
                    else:
                        active_metric = 0.0

                    active_target = float(target_db)
                    metric_label = "dE/E dB"
                else:
                    active_metric = float(current_db)
                    active_target = float(target_db)
                    metric_label = "dB"

                if np.isnan(active_metric) or np.isinf(active_metric):
                    active_metric = 0.0

                self._sim_metric = active_metric
                self._sim_target = active_target

                self.energy_history.append((current_step, active_metric))
                self.energy_history = self.energy_history[-50:]

                if is_step and len(self.linear_energy_history) == 1000 and active_metric <= active_target:
                    if not getattr(self, 'auto_aborted', False):
                        self.auto_aborted = True
                        base_sim_dir = os.path.join(os.path.dirname(getattr(self, 'last_export_path', '')),
                                                    "fdtd_results")
                        if getattr(self, 'total_runs', 1) > 1:
                            c_run = max(1, getattr(self, 'current_run', 1))
                            sim_dir = os.path.join(base_sim_dir, f"run_port_{c_run}")
                        else:
                            sim_dir = base_sim_dir
                        self.sim_thread.graceful_stop(sim_dir)

                time_match = self.re_time.search(line_str)
                speed_match = self.re_speed.search(line_str)

                s_per_ts = 0.0
                elapsed_sec = 0
                total_est_steps = max_ts
                c_run = max(1, getattr(self, 'current_run', 1))
                self._sim_run_str = f" (Run {c_run}/{getattr(self, 'total_runs', 1)})" if getattr(self, 'total_runs',
                                                                                                  1) > 1 else ""

                if time_match and current_step > 0:
                    h = int(time_match.group(1)) if time_match.group(1) else 0
                    m = int(time_match.group(2)) if time_match.group(2) else 0
                    s = int(time_match.group(3)) if time_match.group(3) else 0
                    elapsed_sec = h * 3600 + m * 60 + s
                    s_per_ts = elapsed_sec / current_step
                elif speed_match:
                    s_per_ts = float(speed_match.group(1))

                if s_per_ts > 0:
                    rem_steps_max = max(0, max_ts - current_step)
                    eta_sec_current = rem_steps_max * s_per_ts

                    if len(self.energy_history) >= 10:
                        steps = np.array([pt[0] for pt in self.energy_history])
                        energies = np.array([pt[1] for pt in self.energy_history])
                        slope, intercept = np.polyfit(steps, energies, 1)

                        if slope < -1e-12:
                            rem_steps_energy = (active_target - active_metric) / slope
                            rem_sec_energy = rem_steps_energy * s_per_ts
                            if 0 < rem_sec_energy < eta_sec_current:
                                eta_sec_current = rem_sec_energy
                                total_est_steps = current_step + rem_steps_energy

                        elif slope > 1e-6 and len(self.energy_history) == 50 and current_step > getattr(self,
                                                                                                        'excitation_length',
                                                                                                        0):
                            self._sim_eta_str = "WARNING: Numeric Instability Detected!"
                            eta_sec_current = -1

                    if eta_sec_current >= 0:
                        runs_remaining = getattr(self, 'total_runs', 1) - c_run
                        if getattr(self, 'total_runs', 1) > 1 and len(getattr(self, 'run_durations', [])) > 0:
                            avg_run_time = sum(self.run_durations) / len(self.run_durations)
                            global_eta_sec = eta_sec_current + (runs_remaining * avg_run_time)
                        elif getattr(self, 'total_runs', 1) > 1:
                            estimated_full_run = elapsed_sec + eta_sec_current
                            global_eta_sec = eta_sec_current + (runs_remaining * estimated_full_run)
                        else:
                            global_eta_sec = eta_sec_current

                        total_elapsed_sec = sum(getattr(self, 'run_durations', [])) + (
                            time.time() - self.current_run_start_time if getattr(self, 'current_run_start_time',
                                                                                 None) else elapsed_sec)

                        e_m, e_s = divmod(int(total_elapsed_sec), 60)
                        e_h, e_m = divmod(e_m, 60)
                        r_m, r_s = divmod(int(global_eta_sec), 60)
                        r_h, r_m = divmod(r_m, 60)

                        eta_label = "Global ETA" if getattr(self, 'total_runs', 1) > 1 else "ETA"
                        if e_h > 0 or r_h > 0:
                            self._sim_eta_str = f"Elapsed: {int(e_h)}h {int(e_m)}m | {eta_label}: ~{int(r_h)}h {int(r_m)}m"
                        else:
                            self._sim_eta_str = f"Elapsed: {int(e_m)}m {int(e_s)}s | {eta_label}: ~{int(r_m)}m {int(r_s)}s"

                    if total_est_steps > 0:
                        run_pct = min(1.0, current_step / total_est_steps)
                        prog = int(((c_run - 1) + run_pct) / getattr(self, 'total_runs', 1) * 100)
                        self._sim_progress = min(99, max(1, prog))

                    self._sim_msg = f"Step: {current_step} / {int(total_est_steps)} | Energy: {active_metric:.2f} {metric_label}\n{self._sim_eta_str}\nTarget: {active_target} {metric_label}"

        if found_update:
            latest_payload = {
                "status": "Running",
                "text_status": str(f"Simulation Running{self._sim_run_str}"),
                "bg_color": "#1e3a8a",
                "fg_color": "#bfdbfe",
                "energy_db": float(self._sim_metric),
                "target_db": float(self._sim_target),
                "eta": str(self._sim_eta_str),
                "mesh_cells": str(getattr(self.tab_mesh_settings, 'mesh_cells_str', "Calculated")),
                "progress_pct": int(self._sim_progress)
            }

            wx.CallAfter(self.SetTitle, f"Simulating...{self._sim_run_str} | Energy: {self._sim_metric:.2f} dB")
            write_status_atomic(status_file_path, latest_payload)

        if getattr(self, 'progress_dialog', None) and hasattr(self, '_sim_progress'):
            keep_going = self.progress_dialog.Update(self._sim_progress, self._sim_msg)
            if not keep_going:
                self.sim_thread.abort()
                self.progress_dialog.Destroy()
                self.progress_dialog = None
                wx.CallAfter(self.txt_output.AppendText, "\n[!] Simulation aborted by user.\n")

    def OnSimulationDone(self, success):
        json_path = getattr(self, 'last_export_path', None)
        if getattr(self, 'auto_aborted', False):
            success = True
            self.auto_aborted = False

        if getattr(self, 'progress_dialog', None):
            self.progress_dialog.Destroy()
            self.progress_dialog = None
        self.SetTitle("KiCad Electrodynamic & Static Studio")

        status_val = ""
        status_text = ""
        plot_abs_path = ""

        if getattr(self, 'sim_thread', None) and self.sim_thread.abort_flag:
            status_val = "Interrupted"
            status_text = "Simulation canceled by user."
            self.txt_output.AppendText("\n[!] Simulation aborted by user.\n")
            wx.MessageBox("Simulation aborted.", "Interrupted", wx.OK | wx.ICON_WARNING)
        elif success:
            status_val = "Success"
            status_text = "Simulation completed successfully!"
            sim_dir = os.path.dirname(self.last_export_path)
            plot_abs_path = os.path.join(sim_dir, "fdtd_results", "s_parameter_plot.png")

            self.txt_output.AppendText("\n--- Simulation Complete ---\n")
            self.telegram_bot.send_completion(plot_abs_path)
            wx.MessageBox("Simulation completed successfully!", "Success", wx.OK | wx.ICON_INFORMATION)
        else:
            status_val = "Failed"
            status_text = "Simulation failed. Check logs for details."
            self.txt_output.AppendText("\n[!] Simulation finished with errors or was cancelled.\n")
            wx.MessageBox("Simulation failed!", "Error", wx.OK | wx.ICON_ERROR)

        try:
            with open(status_file_path, "r", encoding="utf-8") as f:
                payload = json.load(f)
        except Exception:
            payload = {}

        payload["status"] = status_val
        payload["text_status"] = status_text
        if plot_abs_path:
            payload["plot_path"] = plot_abs_path

        write_status_atomic(status_file_path, payload)

    def OnLaunchParaView(self, event):
        if not self.last_export_path or not os.path.exists(self.last_export_path):
            wx.MessageBox("Please 'Export to JSON' first.", "Info", wx.OK | wx.ICON_INFORMATION)
            return

        script_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "solvers", "run_openems.py")

        if not os.path.exists(script_path):
            wx.MessageBox(f"Cannot find simulation script at:\n{script_path}", "Error", wx.OK | wx.ICON_ERROR)
            return

        try:
            subprocess.Popen([sys.executable, script_path, self.last_export_path, "--paraview"], creationflags=subprocess.CREATE_NO_WINDOW)
        except Exception as e:
            wx.MessageBox(f"Failed to launch ParaView:\n{e}", "Launch Error", wx.OK | wx.ICON_ERROR)

    def OnClose(self, event):
        self.EndModal(wx.ID_OK)

class EyeDiagramConfigDialog(wx.Dialog):
    def __init__(self, parent):
        super().__init__(parent, title="Eye Diagram Targets", size=(350, 250))
        main_sizer = wx.BoxSizer(wx.VERTICAL)

        grid = wx.FlexGridSizer(3, 2, 10, 10)

        grid.Add(wx.StaticText(self, label="Target Bitrate (bps):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.tc_bitrate = wx.TextCtrl(self, value="5e9")
        grid.Add(self.tc_bitrate, 1, wx.EXPAND)

        grid.Add(wx.StaticText(self, label="Mask Width Ratio (0-1):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.tc_width = wx.TextCtrl(self, value="0.40")
        grid.Add(self.tc_width, 1, wx.EXPAND)

        grid.Add(wx.StaticText(self, label="Mask Height Ratio (0-1):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.tc_height = wx.TextCtrl(self, value="0.40")
        grid.Add(self.tc_height, 1, wx.EXPAND)

        main_sizer.Add(grid, 1, wx.EXPAND | wx.ALL, 15)

        btn_sizer = self.CreateButtonSizer(wx.OK | wx.CANCEL)
        main_sizer.Add(btn_sizer, 0, wx.ALIGN_RIGHT | wx.ALL, 10)
        self.SetSizer(main_sizer)

    def get_values(self):
        return {
            "bitrate": float(self.tc_bitrate.GetValue()),
            "width_ratio": float(self.tc_width.GetValue()),
            "height_ratio": float(self.tc_height.GetValue())
        }


class SPICEZinConfigDialog(wx.Dialog):
    def __init__(self, parent):
        super().__init__(parent, title="SPICE |Zin| Frequency Sweep", size=(350, 200))
        main_sizer = wx.BoxSizer(wx.VERTICAL)

        grid = wx.FlexGridSizer(3, 2, 10, 10)

        grid.Add(wx.StaticText(self, label="Start Frequency (Hz):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.tc_start = wx.TextCtrl(self, value="1e6")
        grid.Add(self.tc_start, 1, wx.EXPAND)

        grid.Add(wx.StaticText(self, label="Stop Frequency (Hz):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.tc_stop = wx.TextCtrl(self, value="10e9")
        grid.Add(self.tc_stop, 1, wx.EXPAND)

        grid.Add(wx.StaticText(self, label="Number of Points:"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.tc_pts = wx.TextCtrl(self, value="1000")
        grid.Add(self.tc_pts, 1, wx.EXPAND)

        main_sizer.Add(grid, 1, wx.EXPAND | wx.ALL, 15)

        btn_sizer = self.CreateButtonSizer(wx.OK | wx.CANCEL)
        main_sizer.Add(btn_sizer, 0, wx.ALIGN_RIGHT | wx.ALL, 10)
        self.SetSizer(main_sizer)

    def get_values(self):
        return {
            "f_start": float(self.tc_start.GetValue()),
            "f_stop": float(self.tc_stop.GetValue()),
            "num_points": int(self.tc_pts.GetValue())
        }


class DCAnalysisTab(wx.Panel):
    def __init__(self, parent, main_app):
        super().__init__(parent)
        self.main_app = main_app
        self.configured_networks = []

        main_sizer = wx.BoxSizer(wx.VERTICAL)

        self.chk_power_only = wx.CheckBox(self, label="Filter Common Power Nets Only (Uncheck to test generic traces)")
        self.chk_power_only.SetValue(True)
        self.chk_power_only.Bind(wx.EVT_CHECKBOX, lambda e: self.refresh_list())
        main_sizer.Add(self.chk_power_only, 0, wx.ALL, 5)

        # Header area
        header_sizer = wx.BoxSizer(wx.HORIZONTAL)
        self.lbl_count = wx.StaticText(self, label="Configured Networks: 0")
        self.lbl_count.SetFont(wx.Font(11, wx.FONTFAMILY_DEFAULT, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_BOLD))
        header_sizer.Add(self.lbl_count, 1, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 5)

        btn_add = wx.Button(self, label="➕ Add DC Network")
        btn_add.SetBackgroundColour(wx.Colour(200, 255, 200))
        btn_add.Bind(wx.EVT_BUTTON, self.on_add_network)
        header_sizer.Add(btn_add, 0, wx.ALL, 5)
        main_sizer.Add(header_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # Scrollable list area
        self.scroll_win = wx.ScrolledWindow(self, style=wx.VSCROLL)
        self.scroll_win.SetScrollRate(5, 5)
        self.list_sizer = wx.BoxSizer(wx.VERTICAL)
        self.scroll_win.SetSizer(self.list_sizer)
        main_sizer.Add(self.scroll_win, 1, wx.EXPAND | wx.ALL, 5)

        # --- Convergence Criteria (Settings Matching the Simulation Tab) ---
        conv_box = wx.StaticBox(self, label=" MNA Convergence Criteria ")
        conv_sizer = wx.StaticBoxSizer(conv_box, wx.VERTICAL)
        grid_conv = wx.FlexGridSizer(2, 2, 8, 12)

        grid_conv.Add(wx.StaticText(conv_box, label="Tolerance:"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_tolerance = wx.TextCtrl(conv_box, value="1e-9")
        grid_conv.Add(self.txt_tolerance, 0, wx.EXPAND)

        grid_conv.Add(wx.StaticText(conv_box, label="Max Iterations:"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_max_iterations = wx.TextCtrl(conv_box, value="5000")
        grid_conv.Add(self.txt_max_iterations, 0, wx.EXPAND)

        grid_conv.AddGrowableCol(1, 1)
        conv_sizer.Add(grid_conv, 1, wx.EXPAND | wx.ALL, 5)
        main_sizer.Add(conv_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # Run Button
        self.btn_run_dc = wx.Button(self, label="Run DC Analysis")
        self.btn_run_dc.SetBackgroundColour(wx.Colour(255, 200, 200))
        self.btn_run_dc.Bind(wx.EVT_BUTTON, self.on_run_dc)
        main_sizer.Add(self.btn_run_dc, 0, wx.EXPAND | wx.ALL, 10)
        self.SetSizer(main_sizer)

    def refresh_list(self):
        self.list_sizer.Clear(True)
        self.lbl_count.SetLabel(f"Configured Networks: {len(self.configured_networks)}")
        for idx, net_data in enumerate(self.configured_networks):
            row_panel = wx.Panel(self.scroll_win)
            row_panel.SetBackgroundColour(wx.Colour(240, 240, 240))
            row_sizer = wx.BoxSizer(wx.HORIZONTAL)

            info_str = f"Network: {net_data['network_name']} | Source: {net_data['vrm_source']['voltage']}V | Load: {net_data['load_sink']['current_A']}A"
            lbl = wx.StaticText(row_panel, label=info_str)
            row_sizer.Add(lbl, 1, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 5)

            btn_edit = wx.Button(row_panel, label="Edit")
            btn_edit.Bind(wx.EVT_BUTTON, lambda evt, i=idx: self.on_edit_network(i))
            row_sizer.Add(btn_edit, 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)

            btn_rem = wx.Button(row_panel, label="Remove")
            btn_rem.SetForegroundColour(wx.RED)
            btn_rem.Bind(wx.EVT_BUTTON, lambda evt, i=idx: self.on_remove_network(i))
            row_sizer.Add(btn_rem, 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)

            row_panel.SetSizer(row_sizer)
            self.list_sizer.Add(row_panel, 0, wx.EXPAND | wx.BOTTOM, 5)

        self.scroll_win.Layout()
        self.scroll_win.Refresh()

    def get_filtered_nets(self):
        all_nets = self.main_app.get_net_names()
        if not self.chk_power_only.IsChecked(): return all_nets
        power_kw = ['gnd', 'ground', 'earth', 'vcc', 'vdd', 'vss', 'vee', 'pwr', 'power', '+', '-', 'v_in', 'v_out',
                    'bus', '0']
        return [n for n in all_nets if any(kw in n.lower() for kw in power_kw) or re.search(r'\d+[vm]v?', n.lower())]

    def on_edit_network(self, idx):
        nets = self.get_filtered_nets()
        dlg = DCNetworkConfigDialog(self, nets, self.main_app.get_pads_for_net, existing_data=self.configured_networks[idx])
        if dlg.ShowModal() == wx.ID_OK:
            self.configured_networks[idx] = dlg.get_network_data()
            self.refresh_list()
        dlg.Destroy()

    def on_add_network(self, event):
        nets = self.get_filtered_nets()
        dlg = DCNetworkConfigDialog(self, nets, self.main_app.get_pads_for_net)
        if dlg.ShowModal() == wx.ID_OK:
            self.configured_networks.append(dlg.get_network_data())
            self.refresh_list()
        dlg.Destroy()

    def on_remove_network(self, idx):
        self.configured_networks.pop(idx)
        self.refresh_list()

    def update_nets(self):
        pass  # Now dynamically handled in the DCNetworkConfigDialog

    def on_run_dc(self, event):
        json_path = getattr(self.main_app, 'last_export_path', None)
        if not json_path or not os.path.exists(json_path):
            wx.MessageBox("Please 'Export to JSON' first.", "Error", wx.OK | wx.ICON_ERROR)
            return

        script_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "solvers", "run_mna.py")
        cmd = [sys.executable, script_path, json_path, "--run-dc"]

        self.btn_run_dc.Disable()
        self.main_app.txt_output.AppendText("\n--- Starting DC Analysis ---\n")

        self.progress_dialog = SimProgressDialog(self.main_app, "DC Analysis")
        self.progress_dialog.lbl_status.SetLabel("Solving MNA Matrix...\n(This is usually very fast)")
        self.progress_dialog.gauge.SetValue(50)
        self.progress_dialog.Show()

        def on_output(chunk_str):
            if getattr(self, 'progress_dialog', None):
                wx.CallAfter(self.progress_dialog.txt_log.AppendText, chunk_str)
                if not self.progress_dialog.Update(50, "Solving MNA Matrix..."):
                    self.sim_thread.abort()
            else:
                wx.CallAfter(self.main_app.txt_output.AppendText, chunk_str)

        def on_done(success):
            if getattr(self, 'progress_dialog', None):
                self.progress_dialog.Destroy()
                self.progress_dialog = None
            self.btn_run_dc.Enable()

            if getattr(self, 'sim_thread', None) and self.sim_thread.abort_flag:
                wx.MessageBox("DC Analysis aborted.", "Interrupted", wx.OK | wx.ICON_WARNING)
            elif success:
                wx.MessageBox("DC Analysis completed successfully!\nHead to Post-Processing.", "Success",
                              wx.OK | wx.ICON_INFORMATION)
            else:
                wx.MessageBox("DC Analysis failed!", "Error", wx.OK | wx.ICON_ERROR)

        self.sim_thread = SimulationThread(cmd, on_output, on_done)
        self.sim_thread.start()

    def get_data(self):
        def safe_float(val, default):
            try:
                return float(val)
            except ValueError:
                return default

        def safe_int(val, default):
            try:
                return int(val)
            except ValueError:
                return default

        return {
            "dc_analysis": {net["network_name"]: net for net in self.configured_networks},
            "dc_convergence": {
                "tolerance": safe_float(self.txt_tolerance.GetValue(), 1e-9),
                "max_iterations": safe_int(self.txt_max_iterations.GetValue(), 5000)
            }
        }

class PostProcessingTab(wx.Panel):
    def __init__(self, parent, main_app):
        super().__init__(parent)
        self.main_app = main_app
        self.current_plot_path = None
        main_sizer = wx.BoxSizer(wx.HORIZONTAL)

        # LEFT PANEL: Controls
        ctrl_panel = wx.ScrolledWindow(self, style=wx.VSCROLL)
        ctrl_panel.SetScrollRate(5, 5)
        ctrl_panel.SetMinSize((400, -1))
        ctrl_sizer = wx.BoxSizer(wx.VERTICAL)

        # 1. SETUP & FILE LOADING
        setup_box = wx.StaticBox(ctrl_panel, label=" 1. Simulation Data Setup ")
        setup_sizer = wx.StaticBoxSizer(setup_box, wx.VERTICAL)

        self.chk_external = wx.CheckBox(setup_box, label="Load different simulation setup (Manual Override)")
        self.chk_external.Bind(wx.EVT_CHECKBOX, self.on_external_toggle)
        setup_sizer.Add(self.chk_external, 0, wx.ALL, 5)

        self.fp_metadata = wx.FilePickerCtrl(setup_box, message="Select simulation_metadata.json",
                                             wildcard="JSON (*.json)|*.json")
        self.fp_metadata.Bind(wx.EVT_FILEPICKER_CHANGED, self.on_metadata_changed)

        self.fp_touchstone = wx.FilePickerCtrl(setup_box, message="Select .sNp", wildcard="Touchstone (*.s*p)|*.s*p")
        self.fp_dc_json = wx.FilePickerCtrl(setup_box, message="Select DC JSON", wildcard="JSON (*.json)|*.json")
        self.fp_vtk = wx.FilePickerCtrl(setup_box, message="Select DC VTK", wildcard="VTK (*.vtk)|*.vtk")

        setup_sizer.Add(wx.StaticText(setup_box, label="Target Metadata JSON (Auto-loads the rest):"), 0,
                        wx.TOP | wx.LEFT, 5)
        setup_sizer.Add(self.fp_metadata, 0, wx.EXPAND | wx.ALL, 5)

        setup_sizer.Add(wx.StaticText(setup_box, label="S-Parameters (.sNp):"), 0, wx.TOP | wx.LEFT, 5)
        setup_sizer.Add(self.fp_touchstone, 0, wx.EXPAND | wx.ALL, 5)

        setup_sizer.Add(wx.StaticText(setup_box, label="DC Output (.json):"), 0, wx.TOP | wx.LEFT, 5)
        setup_sizer.Add(self.fp_dc_json, 0, wx.EXPAND | wx.ALL, 5)

        setup_sizer.Add(wx.StaticText(setup_box, label="DC Potential (.vtk):"), 0, wx.TOP | wx.LEFT, 5)
        setup_sizer.Add(self.fp_vtk, 0, wx.EXPAND | wx.ALL, 5)

        for fp in [self.fp_metadata, self.fp_touchstone, self.fp_dc_json, self.fp_vtk]:
            fp.Enable(False)

        # Network Selection
        hbox_net = wx.BoxSizer(wx.HORIZONTAL)
        self.cmb_dc_network = wx.ComboBox(setup_box, choices=[], style=wx.CB_READONLY)
        self.btn_refresh = wx.Button(setup_box, label="🔄 Refresh")
        self.btn_refresh.Bind(wx.EVT_BUTTON, self.on_metadata_changed)
        hbox_net.Add(wx.StaticText(setup_box, label="Target DC Network:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        hbox_net.Add(self.cmb_dc_network, 1, wx.EXPAND | wx.RIGHT, 5)
        hbox_net.Add(self.btn_refresh, 0, wx.EXPAND)
        setup_sizer.Add(hbox_net, 0, wx.EXPAND | wx.ALL, 5)

        ctrl_sizer.Add(setup_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # 2. SIGNAL INTEGRITY & POWER INTEGRITY
        si_box = wx.StaticBox(ctrl_panel, label=" 2. Signal Integrity and Power Integrity (High-Frequency) ")
        si_sizer = wx.StaticBoxSizer(si_box, wx.VERTICAL)

        # Increased rows to accommodate the 5th button
        si_grid = wx.GridSizer(3, 2, 5, 5)

        btn_sparam = wx.Button(si_box, label="Plot S-Parameters")
        btn_tdr = wx.Button(si_box, label="Plot TDR Impedance")
        btn_xtalk = wx.Button(si_box, label="Plot Crosstalk")
        btn_zin = wx.Button(si_box, label="Plot Input |Zin|")
        btn_spice = wx.Button(si_box, label="Synthesize SPICE Model")
        btn_eye = wx.Button(si_box, label="Plot Eye Diagram (CSV)")

        btn_sparam.Bind(wx.EVT_BUTTON, lambda e: self.dispatch_si("plot_s_params"))
        btn_tdr.Bind(wx.EVT_BUTTON, lambda e: self.dispatch_si("plot_tdr"))
        btn_xtalk.Bind(wx.EVT_BUTTON, lambda e: self.dispatch_si("plot_crosstalk"))
        btn_zin.Bind(wx.EVT_BUTTON, self.on_plot_zin_spice)
        btn_spice.Bind(wx.EVT_BUTTON, self.on_export_spice)
        btn_eye.Bind(wx.EVT_BUTTON, self.on_plot_eye_diagram)

        si_grid.AddMany([btn_sparam, btn_tdr, btn_xtalk, btn_zin, btn_spice, btn_eye])
        si_sizer.Add(si_grid, 0, wx.EXPAND | wx.ALL, 5)
        ctrl_sizer.Add(si_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # 3. DC DIAGNOSTICS
        pi_box = wx.StaticBox(ctrl_panel, label=" 3. DC Analysis Diagnostics ")
        pi_sizer = wx.StaticBoxSizer(pi_box, wx.VERTICAL)

        # Coordinate Target Params (Restored for Heatmaps)
        arg_sizer = wx.FlexGridSizer(2, 4, 5, 5)
        arg_sizer.Add(wx.StaticText(pi_box, label="Target Z (mm):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_z = wx.TextCtrl(pi_box, value="0.0")
        arg_sizer.Add(self.txt_z, 0, wx.EXPAND)

        arg_sizer.Add(wx.StaticText(pi_box, label="Trace Y (mm):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_y = wx.TextCtrl(pi_box, value="10.0")
        arg_sizer.Add(self.txt_y, 0, wx.EXPAND)
        pi_sizer.Add(arg_sizer, 0, wx.EXPAND | wx.ALL, 5)

        pi_grid = wx.GridSizer(2, 2, 5, 5)
        btn_1d = wx.Button(pi_box, label="1D DC Profile")
        btn_2d = wx.Button(pi_box, label="2D Potential Map")
        btn_3d = wx.Button(pi_box, label="3D Surface Render")
        btn_pwr = wx.Button(pi_box, label="Power Density Map")

        btn_1d.Bind(wx.EVT_BUTTON, lambda e: self.dispatch_pi("plot_1d"))
        btn_2d.Bind(wx.EVT_BUTTON, lambda e: self.dispatch_pi("plot_2d"))
        btn_3d.Bind(wx.EVT_BUTTON, lambda e: self.dispatch_pi("plot_3d"))
        btn_pwr.Bind(wx.EVT_BUTTON, lambda e: self.dispatch_pi("plot_power"))

        pi_grid.AddMany([btn_1d, btn_2d, btn_3d, btn_pwr])
        pi_sizer.Add(pi_grid, 0, wx.EXPAND | wx.ALL, 5)
        ctrl_sizer.Add(pi_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # EXPORT CONTROLS
        exp_sizer = wx.BoxSizer(wx.HORIZONTAL)
        btn_add_report = wx.Button(ctrl_panel, label="Add Plot to Report")
        btn_add_report.SetBackgroundColour(wx.Colour(200, 230, 255))
        btn_add_report.Bind(wx.EVT_BUTTON, self.on_add_to_report)
        btn_exp_plot = wx.Button(ctrl_panel, label="Export Current Plot")
        btn_exp_file = wx.Button(ctrl_panel, label="Export Data File")
        btn_exp_plot.Bind(wx.EVT_BUTTON, self.on_export_plot)
        btn_exp_file.Bind(wx.EVT_BUTTON, self.on_export_file)
        exp_sizer.Add(btn_add_report, 1, wx.ALL, 5)
        exp_sizer.Add(btn_exp_plot, 1, wx.ALL, 5)
        exp_sizer.Add(btn_exp_file, 1, wx.ALL, 5)
        ctrl_sizer.Add(exp_sizer, 0, wx.EXPAND | wx.ALL, 5)

        ctrl_panel.SetSizer(ctrl_sizer)
        main_sizer.Add(ctrl_panel, 0, wx.EXPAND | wx.ALL, 5)

        # RIGHT PANEL: Image Viewer
        self.img_panel = wx.ScrolledWindow(self, style=wx.VSCROLL | wx.HSCROLL)
        self.img_panel.SetBackgroundColour(wx.Colour(255, 255, 255))
        self.img_sizer = wx.BoxSizer(wx.VERTICAL)
        self.img_bitmap = wx.StaticBitmap(self.img_panel, wx.ID_ANY, wx.NullBitmap)
        self.img_sizer.Add(self.img_bitmap, 1, wx.ALIGN_CENTER | wx.ALL, 10)
        self.img_panel.SetSizer(self.img_sizer)
        main_sizer.Add(self.img_panel, 1, wx.EXPAND | wx.ALL, 5)

        self.SetSizer(main_sizer)

    def on_external_toggle(self, event):
        enabled = self.chk_external.IsChecked()
        for fp in [self.fp_metadata, self.fp_touchstone, self.fp_dc_json, self.fp_vtk]:
            fp.Enable(enabled)
        self.on_metadata_changed()

    def get_paths(self):
        json_path = self.fp_metadata.GetPath() if self.chk_external.IsChecked() else getattr(self.main_app,
                                                                                             'last_export_path', None)
        if not json_path or not os.path.exists(json_path):
            return "", "", ""

        try:
            import json
            with open(json_path, 'r', encoding='utf-8') as f:
                meta = json.load(f)
            fdtd_dir = meta.get("simulation_dir", os.path.join(os.path.dirname(json_path), "fdtd_results"))
            dc_dir = meta.get("dc_simulation_dir", os.path.join(os.path.dirname(json_path), "dc_results"))
        except Exception:
            fdtd_dir = os.path.join(os.path.dirname(json_path), "fdtd_results")
            dc_dir = os.path.join(os.path.dirname(json_path), "dc_results")

        sp_path, dc_json, vtk_path = "", "", ""

        # S-Parameters are in the FDTD directory
        if os.path.exists(fdtd_dir):
            found_sp_files = []
            for f in os.listdir(fdtd_dir):
                if f.endswith(".s1p") or f.endswith(".s2p") or f.endswith(".s4p") or f.endswith(".sNp"):
                    found_sp_files.append(os.path.join(fdtd_dir, f))

            if found_sp_files:
                found_sp_files.sort(key=os.path.getmtime, reverse=True)
                sp_path = found_sp_files[0]

        if os.path.exists(dc_dir):
            for f in os.listdir(dc_dir):
                if f == "dc_ir_drop.json":
                    dc_json = os.path.join(dc_dir, f)
                if f.endswith(".vtk") and "potential" in f.lower():
                    vtk_path = os.path.join(dc_dir, f)

        return sp_path, dc_json, vtk_path

    def on_metadata_changed(self, event=None):
        json_path = self.fp_metadata.GetPath() if self.chk_external.IsChecked() else getattr(self.main_app,
                                                                                             'last_export_path', None)

        if not json_path or not os.path.exists(json_path):
            self.cmb_dc_network.Clear()
            return

        try:
            import json
            with open(json_path, 'r', encoding='utf-8') as f:
                data = json.load(f)

            networks = list(data.get("dc_analysis", {}).keys())
            self.cmb_dc_network.SetItems(networks)
            if networks:
                self.cmb_dc_network.SetSelection(0)
            else:
                self.cmb_dc_network.Clear()

            fdtd_dir = data.get("simulation_dir", os.path.join(os.path.dirname(json_path), "fdtd_results"))
            dc_dir = data.get("dc_simulation_dir", os.path.join(os.path.dirname(json_path), "dc_results"))

            if os.path.exists(fdtd_dir):
                for f in os.listdir(fdtd_dir):
                    if f.endswith(".s2p") or f.endswith(".s4p") or f.endswith(".sNp"):
                        self.fp_touchstone.SetPath(os.path.join(fdtd_dir, f))

            if os.path.exists(dc_dir):
                for f in os.listdir(dc_dir):
                    if f == "dc_ir_drop.json":
                        self.fp_dc_json.SetPath(os.path.join(dc_dir, f))
                    elif f.endswith(".vtk") and "potential" in f.lower():
                        self.fp_vtk.SetPath(os.path.join(dc_dir, f))

        except Exception as e:
            wx.MessageBox(f"Failed to load networks from JSON: {e}", "Error", wx.OK | wx.ICON_ERROR)

    def on_add_to_report(self, event):
        if not self.current_plot_path or not os.path.exists(self.current_plot_path):
            wx.MessageBox("No plot is currently generated or displayed.", "Info", wx.OK | wx.ICON_INFORMATION)
            return

        self.main_app.tab_report.add_image(self.current_plot_path)
        wx.MessageBox("Plot successfully added to the report queue!", "Added", wx.OK | wx.ICON_INFORMATION)

    def on_plot_zin_spice(self, event):
        with wx.FileDialog(self, "Select SPICE Transient CSV Export", wildcard="CSV Files (*.csv)|*.csv",
                           style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST) as fileDialog:
            if fileDialog.ShowModal() == wx.ID_CANCEL:
                return
            csv_path = fileDialog.GetPath()

        dlg = SPICEZinConfigDialog(self)
        if dlg.ShowModal() == wx.ID_OK:
            params = dlg.get_values()
            wx.BeginBusyCursor()
            try:
                out_png = csv_path.replace(os.path.splitext(csv_path)[1], "_zin.png")
                post_processing.plot_spice_impedance_magnitude(
                    csv_path,
                    f_start=params["f_start"],
                    f_stop=params["f_stop"],
                    num_points=params["num_points"]
                )
                if os.path.exists(out_png):
                    self.current_plot_path = out_png
                    self.display_image(out_png)
            except Exception as e:
                wx.MessageBox(f"|Zin| Plotting Error: {e}", "Error", wx.OK | wx.ICON_ERROR)
            wx.EndBusyCursor()
        dlg.Destroy()

    def on_export_spice(self, event):
        sp_path, dc_json, _ = self.get_paths()
        if not sp_path or not os.path.exists(sp_path):
            wx.MessageBox("S-Parameter output not found in simulation folder.", "Error", wx.OK | wx.ICON_ERROR)
            return

        include_dc = False
        if dc_json and os.path.exists(dc_json):
            resp = wx.MessageBox("DC Analysis data found. Stitch DC behavior into the SPICE model?", "Include DC?",
                                 wx.YES_NO | wx.ICON_QUESTION)
            include_dc = (resp == wx.YES)

        base_name = os.path.splitext(os.path.basename(sp_path))[0]
        default_file = f"{base_name}_model.cir"

        with wx.FileDialog(self, "Save SPICE Model", defaultFile=default_file,
                           wildcard="SPICE Circuit (*.cir)|*.cir|SPICE Library (*.lib)|*.lib",
                           style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT) as dlg:
            if dlg.ShowModal() == wx.ID_CANCEL: return
            out_path = dlg.GetPath()
            wx.BeginBusyCursor()
            try:
                post_processing.convert_s_params_to_spice(sp_path, dc_json if include_dc else None, out_path)
                wx.MessageBox(f"SPICE Model saved to {out_path}", "Success", wx.OK)
            except Exception as e:
                wx.MessageBox(f"SPICE Synthesis failed: {e}", "Error", wx.OK | wx.ICON_ERROR)
            wx.EndBusyCursor()

    def on_plot_eye_diagram(self, event):
        with wx.FileDialog(self, "Select SPICE Transient CSV Export", wildcard="CSV Files (*.csv)|*.csv",
                           style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST) as fileDialog:
            if fileDialog.ShowModal() == wx.ID_CANCEL:
                return
            csv_path = fileDialog.GetPath()

        dlg = EyeDiagramConfigDialog(self)
        if dlg.ShowModal() == wx.ID_OK:
            params = dlg.get_values()
            wx.BeginBusyCursor()
            try:
                out_png = csv_path.replace(os.path.splitext(csv_path)[1], "_eye_diagram.png")
                post_processing.plot_eye_diagram_with_mask(
                    csv_path,
                    bitrate_bps=params["bitrate"],
                    mask_width_ratio=params["width_ratio"],
                    mask_height_ratio=params["height_ratio"],
                    out_path=out_png
                )

                if os.path.exists(out_png):
                    self.current_plot_path = out_png
                    self.display_image(out_png)
            except Exception as e:
                wx.MessageBox(f"Eye Diagram Plotting Error: {e}", "Error", wx.OK | wx.ICON_ERROR)
            wx.EndBusyCursor()
        dlg.Destroy()

    def dispatch_si(self, task):
        sp_path, _, _ = self.get_paths()
        if not sp_path or not os.path.exists(sp_path):
            wx.MessageBox("S-Parameter file not found in simulation folder.", "Error", wx.OK | wx.ICON_ERROR)
            return

        wx.BeginBusyCursor()
        try:
            out_png = sp_path.replace(os.path.splitext(sp_path)[1], f"_{task}.png")
            if task == "plot_s_params":
                post_processing.plot_touchstone(sp_path)
            elif task == "plot_tdr":
                post_processing.export_tdr_impedance(sp_path)
            elif task == "plot_crosstalk":
                post_processing.export_crosstalk(sp_path)

            if os.path.exists(out_png):
                self.current_plot_path = out_png
                self.display_image(out_png)
        except Exception as e:
            wx.MessageBox(f"SI Error: {e}", "Error", wx.OK | wx.ICON_ERROR)
        wx.EndBusyCursor()

    def dispatch_pi(self, task):
        _, dc_json, vtk_path = self.get_paths()
        if not vtk_path or not os.path.exists(vtk_path):
            wx.MessageBox("DC VTK file not found. Ensure DC Analysis has been run.", "Error", wx.OK | wx.ICON_ERROR)
            return

        if not self.cmb_dc_network.GetValue():
            wx.MessageBox("Please select a valid DC Network to analyze.", "Error", wx.OK | wx.ICON_ERROR)
            return

        network_name = self.cmb_dc_network.GetValue()
        json_path = self.fp_metadata.GetPath() if self.chk_external.IsChecked() else getattr(self.main_app,
                                                                                             'last_export_path', None)

        wx.BeginBusyCursor()
        try:


            out_png = vtk_path.replace(os.path.splitext(vtk_path)[1], f"_{task}.png")

            pot_mesh = pv.read(vtk_path)
            x_coords = np.array(pot_mesh.x)
            y_coords = np.array(pot_mesh.y)
            z_coords = np.array(pot_mesh.z)
            dims = (len(x_coords), len(y_coords), len(z_coords))

            pot_key = next((k for k in pot_mesh.point_data.keys() if 'potential' in k.lower()), None)
            if not pot_key: raise KeyError("Potential array missing.")
            potential_3d = pot_mesh.point_data[pot_key].reshape(dims, order='F')

            cond_path = vtk_path.replace("potential", "conductivity")
            sigma_map = None
            if os.path.exists(cond_path):
                cond_mesh = pv.read(cond_path)
                cond_key = next((k for k in cond_mesh.point_data.keys() if 'conductivity' in k.lower()), None)
                if cond_key:
                    sigma_map = cond_mesh.point_data[cond_key].reshape(dims, order='F')

            try:
                z_val = float(self.txt_z.GetValue())
                y_val = float(self.txt_y.GetValue())
            except ValueError:
                z_val, y_val = 0.0, 10.0

            if task == "plot_1d":
                plotter.plot_dc_ir_drop_profile(potential_3d, x_coords, y_coords, z_coords, sigma_map, z_val, y_val,
                                                json_path, network_name, out_png)
            elif task == "plot_2d":
                plotter.plot_potential_map(potential_3d, x_coords, y_coords, z_coords, sigma_map, z_val, out_png)
            elif task == "plot_3d":
                plotter.plot_3d_potential(potential_3d, x_coords, y_coords, z_coords, sigma_map, out_png)
            elif task == "plot_power":
                plotter.plot_power_density(potential_3d, x_coords, y_coords, z_coords, sigma_map, z_val, 58e6, out_png)

            if os.path.exists(out_png):
                self.current_plot_path = out_png
                self.display_image(out_png)
        except Exception as e:
            wx.MessageBox(f"PI Analysis Error: {e}", "Error", wx.OK | wx.ICON_ERROR)
        finally:
            wx.EndBusyCursor()

    def on_export_plot(self, event):
        if not self.current_plot_path or not os.path.exists(self.current_plot_path): return
        with wx.FileDialog(self, "Export Plot", wildcard="PNG (*.png)|*.png",
                           style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT) as dlg:
            if dlg.ShowModal() == wx.ID_OK:
                import shutil
                shutil.copy(self.current_plot_path, dlg.GetPath())

    def on_export_file(self, event):
        sp_path, dc_json, vtk_path = self.get_paths()
        if not sp_path and not dc_json: return
        with wx.FileDialog(self, "Export Primary Data File", style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT) as dlg:
            if dlg.ShowModal() == wx.ID_OK:
                import shutil
                shutil.copy(sp_path if sp_path else dc_json, dlg.GetPath())

    def display_image(self, filename):
        if os.path.exists(filename):
            self.current_plot_path = filename
            img = wx.Image(filename, wx.BITMAP_TYPE_PNG)
            panel_w, panel_h = self.img_panel.GetSize()
            if panel_w > 50 and panel_h > 50:
                ratio = min(panel_w / float(img.GetWidth()), panel_h / float(img.GetHeight()))
                if ratio < 1.0: img = img.Scale(int(img.GetWidth() * ratio), int(img.GetHeight() * ratio),
                                                wx.IMAGE_QUALITY_HIGH)
            self.img_bitmap.SetBitmap(wx.Bitmap(img))
            self.img_panel.Layout()

class ReportGeneratorTab(wx.Panel):
    def __init__(self, parent, main_app):
        super().__init__(parent)
        self.main_app = main_app
        self.image_queue = []

        main_sizer = wx.BoxSizer(wx.VERTICAL)

        # Header
        header_sizer = wx.BoxSizer(wx.HORIZONTAL)
        lbl_title = wx.StaticText(self, label="Report Images Queue:")
        lbl_title.SetFont(wx.Font(11, wx.FONTFAMILY_DEFAULT, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_BOLD))
        header_sizer.Add(lbl_title, 1, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 5)
        main_sizer.Add(header_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # ListBox for queued images
        self.list_images = wx.ListBox(self, style=wx.LB_EXTENDED)
        main_sizer.Add(self.list_images, 1, wx.EXPAND | wx.ALL, 5)

        # Control Buttons
        btn_sizer = wx.BoxSizer(wx.HORIZONTAL)
        btn_add_file = wx.Button(self, label="Add Photo from File")
        btn_add_clip = wx.Button(self, label="Add Photo from Clipboard")
        btn_remove = wx.Button(self, label="Remove Selected")
        btn_generate = wx.Button(self, label="Generate PPTX Report")

        btn_generate.SetBackgroundColour(wx.Colour(255, 220, 180))
        btn_generate.SetFont(wx.Font(10, wx.FONTFAMILY_DEFAULT, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_BOLD))

        btn_add_file.Bind(wx.EVT_BUTTON, self.on_add_file)
        btn_add_clip.Bind(wx.EVT_BUTTON, self.on_add_clip)
        btn_remove.Bind(wx.EVT_BUTTON, self.on_remove)
        btn_generate.Bind(wx.EVT_BUTTON, self.on_generate)

        btn_sizer.Add(btn_add_file, 0, wx.ALL, 5)
        btn_sizer.Add(btn_add_clip, 0, wx.ALL, 5)
        btn_sizer.Add(btn_remove, 0, wx.ALL, 5)
        btn_sizer.AddStretchSpacer()
        btn_sizer.Add(btn_generate, 0, wx.ALL, 5)

        main_sizer.Add(btn_sizer, 0, wx.EXPAND | wx.ALL, 5)
        self.SetSizer(main_sizer)

    def add_image(self, filepath):
        self.image_queue.append(filepath)
        self.list_images.Append(os.path.basename(filepath))

    def on_add_file(self, event):
        with wx.FileDialog(self, "Select Image", wildcard="Image files (*.png;*.jpg;*.jpeg)|*.png;*.jpg;*.jpeg",
                           style=wx.FD_OPEN | wx.FD_MULTIPLE) as dlg:
            if dlg.ShowModal() == wx.ID_OK:
                for path in dlg.GetPaths():
                    self.add_image(path)

    def on_add_clip(self, event):
        if wx.TheClipboard.Open():
            if wx.TheClipboard.IsSupported(wx.DataFormat(wx.DF_BITMAP)):
                data = wx.BitmapDataObject()
                wx.TheClipboard.GetData(data)
                wx.TheClipboard.Close()

                bmp = data.GetBitmap()
                img = bmp.ConvertToImage()
                temp_path = os.path.join(tempfile.gettempdir(), f"clip_{int(time.time())}.png")
                img.SaveFile(temp_path, wx.BITMAP_TYPE_PNG)
                self.add_image(temp_path)
            else:
                wx.TheClipboard.Close()
                wx.MessageBox("No valid image found in clipboard.", "Warning", wx.OK | wx.ICON_WARNING)

    def on_remove(self, event):
        selections = self.list_images.GetSelections()
        for i in reversed(selections):
            self.list_images.Delete(i)
            del self.image_queue[i]

    def on_generate(self, event):
        if not self.image_queue:
            wx.MessageBox("No images added to the report.", "Warning", wx.OK | wx.ICON_WARNING)
            return

        with wx.FileDialog(self, "Save Presentation", wildcard="PowerPoint (*.pptx)|*.pptx",
                           defaultFile="Simulation_Report.pptx", style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT) as dlg:
            if dlg.ShowModal() == wx.ID_CANCEL:
                return
            out_path = dlg.GetPath()

        wx.BeginBusyCursor()
        try:
            self._generate_pptx(out_path)
        finally:
            wx.EndBusyCursor()

    def _generate_pptx(self, out_path):
        template_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "resources", "template.pptx")

        if os.path.exists(template_path):
            prs = Presentation(template_path)
            title_layout = prs.slide_layouts[0]
            # Assumes layout 1 is Content, adjust index if your specific template differs
            content_layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[6]
        else:
            prs = Presentation()
            title_layout = prs.slide_layouts[0]
            content_layout = prs.slide_layouts[6]  # Blank fallback

        # 1. Main Title Slide
        slide = prs.slides.add_slide(title_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1] if len(slide.placeholders) > 1 else None

        project_name = self.main_app.board.get_project().name if self.main_app.board else "KiCad PCB"
        if title:
            title.text = f"{project_name} - Electrodynamic & Static Analysis"

        if subtitle:
            timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
            subtitle.text = f"Generated on {timestamp}"

        for img_path in self.image_queue:
            if not os.path.exists(img_path):
                continue

            slide = prs.slides.add_slide(content_layout)

            try:
                pic = slide.shapes.add_picture(img_path, 0, 0, width=Inches(7.5))
                pic.left = int((prs.slide_width - pic.width) / 2)
                pic.top = int((prs.slide_height - pic.height) / 2) + int(Inches(0.4))
            except Exception as e:
                print(f"[!] Failed to insert image {img_path}: {e}")

        try:
            prs.save(out_path)
            wx.MessageBox(f"Report successfully generated at:\n{out_path}", "Success", wx.OK | wx.ICON_INFORMATION)
        except Exception as e:
            wx.MessageBox(f"Failed to save report: {e}", "Error", wx.OK | wx.ICON_ERROR)

class SettingsTab(wx.Panel):
    def __init__(self, parent):
        super().__init__(parent)
        main_sizer = wx.BoxSizer(wx.VERTICAL)

        self.btn_scan = wx.Button(self, label="Scan System Dependencies")
        self.btn_scan.Bind(wx.EVT_BUTTON, self.on_scan)
        main_sizer.Add(self.btn_scan, 0, wx.EXPAND | wx.ALL, 10)

        self.txt_log = wx.TextCtrl(self, style=wx.TE_MULTILINE | wx.TE_READONLY)
        main_sizer.Add(self.txt_log, 1, wx.EXPAND | wx.ALL, 5)
        self.SetSizer(main_sizer)

    def on_scan(self, event):
        self.txt_log.Clear()

        # Check openEMS Native
        openems_path = shutil.which("openEMS")
        self.txt_log.AppendText(f"[+] openEMS Binary: {'Found' if openems_path else 'Missing'}\n")

        if openems_path:
            openems_locator(self.txt_log)

        # Check ParaView
        paraview_path = shutil.which("paraview")
        self.txt_log.AppendText(f"[+] ParaView: {'Found' if paraview_path else 'Missing'}\n")

        # Check WSL
        try:
            startupinfo = subprocess.STARTUPINFO()
            startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
            creationflags = subprocess.CREATE_NO_WINDOW
            wsl_check = subprocess.run(["wsl", "-l", "-v"], capture_output=True, text=True)
            if wsl_check.returncode == 0:
                self.txt_log.AppendText("[+] WSL: Installed\n")
                # Check Python inside WSL
                wsl_py = subprocess.run(["wsl", "python3", "-c", r'"import openEMS"'], capture_output=True)
                self.txt_log.AppendText(
                    f"  -> openEMS Python Module (WSL): {'Found' if wsl_py.returncode == 0 else 'Missing'}\n")
        except FileNotFoundError:
            self.txt_log.AppendText("[-] WSL: Not Found\n")

class NetManagerTab(wx.Panel):
    def __init__(self, parent, board):
        super().__init__(parent)
        self.board = board
        self.all_nets = []
        self.filtered_nets = []

        main_sizer = wx.BoxSizer(wx.VERTICAL)

        # Search & Filter Bar
        hbox_search = wx.BoxSizer(wx.HORIZONTAL)
        hbox_search.Add(wx.StaticText(self, label="Search/Filter Nets (e.g. GND, VCC):"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.search_ctrl = wx.SearchCtrl(self, style=wx.TE_PROCESS_ENTER)
        self.search_ctrl.ShowSearchButton(True)
        self.search_ctrl.ShowCancelButton(True)
        self.search_ctrl.Bind(wx.EVT_TEXT, self.OnSearch)
        self.search_ctrl.Bind(wx.EVT_SEARCHCTRL_CANCEL_BTN, self.OnCancelSearch)
        hbox_search.Add(self.search_ctrl, 1, wx.EXPAND | wx.ALL, 5)
        main_sizer.Add(hbox_search, 0, wx.EXPAND | wx.ALL, 5)

        # Selection Controls
        hbox_btns = wx.BoxSizer(wx.HORIZONTAL)
        btn_sel_all = wx.Button(self, label="Select All")
        btn_desel_all = wx.Button(self, label="Deselect All")
        btn_sel_all.Bind(wx.EVT_BUTTON, self.OnSelectAll)
        btn_desel_all.Bind(wx.EVT_BUTTON, self.OnDeselectAll)
        hbox_btns.Add(btn_sel_all, 0, wx.RIGHT, 5)
        hbox_btns.Add(btn_desel_all, 0, wx.LEFT, 5)
        main_sizer.Add(hbox_btns, 0, wx.EXPAND | wx.ALL, 5)

        # Net Checklist
        self.check_list = wx.CheckListBox(self, choices=[], style=wx.LB_EXTENDED | wx.LB_HSCROLL)
        main_sizer.Add(self.check_list, 1, wx.EXPAND | wx.ALL, 5)

        # Log Output Console
        self.txt_output = wx.TextCtrl(self, style=wx.TE_MULTILINE | wx.TE_READONLY)
        self.txt_output.SetFont(wx.Font(10, wx.FONTFAMILY_TELETYPE, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_NORMAL))
        main_sizer.Add(self.txt_output, 1, wx.EXPAND | wx.ALL, 5)

        self.SetSizer(main_sizer)
        self.LoadNets()

    def LoadNets(self):
        try:
            nets = self.board.get_nets()
            self.all_nets = sorted([net.name for net in nets if net.name])
        except Exception as e:
            self.all_nets = [f"Error loading nets: {e}"]
        self.filtered_nets = list(self.all_nets)
        self.RefreshListUI()

    def RefreshListUI(self):
        checked_nets = self.GetSelectedNets()
        self.check_list.Clear()
        for net_name in self.filtered_nets:
            idx = self.check_list.Append(net_name)
            if net_name in checked_nets:
                self.check_list.Check(idx, True)

    def GetSelectedNets(self):
        nets = []
        for i in range(self.check_list.GetCount()):
            if self.check_list.IsChecked(i):
                nets.append(self.check_list.GetString(i))
        return nets

    def OnSearch(self, event):
        query = self.search_ctrl.GetValue().lower()
        if not query:
            self.filtered_nets = list(self.all_nets)
        else:
            self.filtered_nets = [net for net in self.all_nets if query in net.lower()]
        self.RefreshListUI()

    def OnCancelSearch(self, event):
        self.search_ctrl.SetValue("")
        self.OnSearch(None)

    def OnSelectAll(self, event):
        for i in range(self.check_list.GetCount()):
            self.check_list.Check(i, True)

    def OnDeselectAll(self, event):
        for i in range(self.check_list.GetCount()):
            self.check_list.Check(i, False)

class ComponentManagerTab(wx.Panel):
    def __init__(self, parent, board):
        super().__init__(parent)
        self.board = board
        self.all_components = []
        self.filtered_components = []

        main_sizer = wx.BoxSizer(wx.VERTICAL)

        # --- Search & Filter Bar ---
        hbox_search = wx.BoxSizer(wx.HORIZONTAL)
        hbox_search.Add(wx.StaticText(self, label="Search/Filter (e.g. R1, 10k):"), 0,
                        wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.search_ctrl = wx.SearchCtrl(self, style=wx.TE_PROCESS_ENTER)
        self.search_ctrl.ShowSearchButton(True)
        self.search_ctrl.ShowCancelButton(True)
        self.search_ctrl.Bind(wx.EVT_TEXT, self.OnSearch)
        self.search_ctrl.Bind(wx.EVT_SEARCHCTRL_CANCEL_BTN, self.OnCancelSearch)
        hbox_search.Add(self.search_ctrl, 1, wx.EXPAND | wx.ALL, 5)
        main_sizer.Add(hbox_search, 0, wx.EXPAND | wx.ALL, 5)

        # --- Selection Controls ---
        hbox_btns = wx.BoxSizer(wx.HORIZONTAL)
        btn_sel_all = wx.Button(self, label="Select All")
        btn_desel_all = wx.Button(self, label="Deselect All")
        btn_sel_all.Bind(wx.EVT_BUTTON, self.OnSelectAll)
        btn_desel_all.Bind(wx.EVT_BUTTON, self.OnDeselectAll)
        hbox_btns.Add(btn_sel_all, 0, wx.RIGHT, 5)
        hbox_btns.Add(btn_desel_all, 0, wx.LEFT, 5)
        main_sizer.Add(hbox_btns, 0, wx.EXPAND | wx.ALL, 5)

        # --- Component Checklist ---
        self.check_list = wx.CheckListBox(self, choices=[], style=wx.LB_EXTENDED | wx.LB_HSCROLL)
        self.check_list.Bind(wx.EVT_RIGHT_DOWN, self.OnRightClick)
        main_sizer.Add(self.check_list, 1, wx.EXPAND | wx.ALL, 5)

        self.SetSizer(main_sizer)
        self.LoadComponents()

    def LoadComponents(self):
        """Extract valid RLC components, ignoring connectors and raw pads."""
        self.all_components = []
        try:
            for fp in self.board.get_footprints():
                ref = "Unknown"
                val = ""
                if hasattr(fp, 'reference_field') and fp.reference_field:
                    ref = fp.reference_field.text.value if hasattr(fp.reference_field.text, 'value') else str(
                        fp.reference_field.text)
                elif hasattr(fp, 'reference'):
                    ref = str(fp.reference)

                if hasattr(fp, 'value_field') and fp.value_field:
                    val = fp.value_field.text.value if hasattr(fp.value_field.text, 'value') else str(
                        fp.value_field.text)

                # Filter R, L, C, U, FB (Exclude J, P connectors)
                if not ref.startswith(('R', 'L', 'C', 'U', 'FB', 'L')):
                    continue

                # Fetch pads to calculate the physical bridging box
                pads = getattr(fp, 'pads', None)
                if pads is None and hasattr(fp, 'definition'):
                    pads = getattr(fp.definition, 'pads', [])

                if len(pads) < 2:
                    continue  # Need at least 2 pads for a lumped element

                # Store raw KiCAD data
                comp_data = {
                    "reference": ref,
                    "value": val,
                    "layer": fp.layer if hasattr(fp, 'layer') else "F.Cu",
                    "pads": [{"x": p.position.x / 1e6, "y": p.position.y / 1e6} for p in pads[:2]]
                    # Take first 2 pads for bridging
                }

                display_str = f"{ref}  |  Value: {val}  |  Layer: {comp_data['layer']}"
                self.all_components.append((display_str, comp_data))

        except Exception as e:
            print(f"Error loading components: {e}")

        self.filtered_components = list(self.all_components)
        self.RefreshListUI()

    def RefreshListUI(self):
        checked_refs = self.GetSelectedReferences()
        self.check_list.Clear()
        for display_str, data in self.filtered_components:
            idx = self.check_list.Append(display_str, data)
            if data["reference"] in checked_refs:
                self.check_list.Check(idx, True)

    def GetSelectedReferences(self):
        """Keep track of what is checked even during filtering."""
        refs = []
        for i in range(self.check_list.GetCount()):
            if self.check_list.IsChecked(i):
                data = self.check_list.GetClientData(i)
                refs.append(data["reference"])
        return refs

    def OnSearch(self, event):
        query = self.search_ctrl.GetValue().lower()
        if not query:
            self.filtered_components = list(self.all_components)
        else:
            self.filtered_components = [
                item for item in self.all_components
                if query in item[0].lower()
            ]
        self.RefreshListUI()

    def OnRightClick(self, event):
        idx = self.check_list.HitTest(event.GetPosition())
        if idx != wx.NOT_FOUND:
            self.check_list.SetSelection(idx)
            data = self.check_list.GetClientData(idx)
            ref = data.get("reference", "")

            # Only allow RLC components to have parasitics edited
            if ref.startswith(('R', 'L', 'C')):
                dlg = ParasiticsDialog(self, data)
                if dlg.ShowModal() == wx.ID_OK:
                    data.update(dlg.get_values())
                    self.check_list.SetClientData(idx, data)
                dlg.Destroy()
        event.Skip()

    def OnCancelSearch(self, event):
        self.search_ctrl.SetValue("")
        self.OnSearch(None)

    def OnSelectAll(self, event):
        for i in range(self.check_list.GetCount()):
            self.check_list.Check(i, True)

    def OnDeselectAll(self, event):
        for i in range(self.check_list.GetCount()):
            self.check_list.Check(i, False)

    def get_data(self):
        """Returns JSON exportable data for selected components."""
        selected = []
        for i in range(self.check_list.GetCount()):
            if self.check_list.IsChecked(i):
                selected.append(self.check_list.GetClientData(i))
        return {"discrete_components": selected}

class SimSettingsTab(wx.Panel):
    def __init__(self, parent):
        super().__init__(parent)
        main_sizer = wx.BoxSizer(wx.VERTICAL)

        freq_box = wx.StaticBox(self, label=" Frequency & Excitation Sweep ")
        freq_sizer = wx.StaticBoxSizer(freq_box, wx.VERTICAL)
        grid_freq = wx.FlexGridSizer(5, 2, 8, 12)

        grid_freq.Add(wx.StaticText(freq_box, label="Excitation Type:"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.cmb_excite_type = wx.ComboBox(freq_box, choices=["Gaussian Pulse", "Sinusoid", "Step", "Dirac Impulse"], style=wx.CB_READONLY)
        self.cmb_excite_type.SetValue("Gaussian Pulse")
        grid_freq.Add(self.cmb_excite_type, 0, wx.EXPAND)

        grid_freq.Add(wx.StaticText(freq_box, label="Start / Sine Frequency (GHz):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_f_start = wx.TextCtrl(freq_box, value="0.1")
        grid_freq.Add(self.txt_f_start, 0, wx.EXPAND)

        grid_freq.Add(wx.StaticText(freq_box, label="Stop Frequency (GHz):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_f_stop = wx.TextCtrl(freq_box, value="10.0")
        grid_freq.Add(self.txt_f_stop, 0, wx.EXPAND)

        grid_freq.Add(wx.StaticText(freq_box, label="Pulse Max Frequency (f_max GHz):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_f_max = wx.TextCtrl(freq_box, value="12.0")
        grid_freq.Add(self.txt_f_max, 0, wx.EXPAND)

        grid_freq.Add(wx.StaticText(freq_box, label="Frequency Sweep Points:"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_num_points = wx.TextCtrl(freq_box, value="1001")
        grid_freq.Add(self.txt_num_points, 0, wx.EXPAND)

        grid_freq.AddGrowableCol(1, 1)
        freq_sizer.Add(grid_freq, 1, wx.EXPAND | wx.ALL, 8)
        main_sizer.Add(freq_sizer, 0, wx.EXPAND | wx.ALL, 10)

        # --- Convergence Criteria ---
        conv_box = wx.StaticBox(self, label=" Convergence Criteria ")
        conv_sizer = wx.StaticBoxSizer(conv_box, wx.VERTICAL)
        grid_conv = wx.FlexGridSizer(0, 2, 8, 12)

        grid_conv.Add(wx.StaticText(conv_box, label="Energy Limit Target (dB):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.cmb_energy_limit = wx.TextCtrl(conv_box, value="-40.0")
        self.cmb_energy_limit.SetValue("-40.0")
        grid_conv.Add(self.cmb_energy_limit, 0, wx.EXPAND)

        grid_conv.Add(wx.StaticText(conv_box, label="Max Timesteps Limit:"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_max_timesteps = wx.TextCtrl(conv_box, value="100000")
        grid_conv.Add(self.txt_max_timesteps, 0, wx.EXPAND)

        # --- AR Filter UI ---
        self.chk_ar_filter = wx.CheckBox(conv_box, label="Use AR Filter (Fix 0Hz Drop)")
        self.chk_ar_filter.SetValue(False)
        self.chk_ar_filter.SetToolTip(
            "Extrapolates the time-domain signal to infinity to calculate an accurate 0Hz/DC response.")
        grid_conv.Add(self.chk_ar_filter, 0, wx.ALIGN_CENTER_VERTICAL)

        hbox_ar = wx.BoxSizer(wx.HORIZONTAL)
        hbox_ar.Add(wx.StaticText(conv_box, label="AR Extrap. Steps:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.txt_ar_steps = wx.TextCtrl(conv_box, value="100000")
        hbox_ar.Add(self.txt_ar_steps, 1, wx.EXPAND)
        grid_conv.Add(hbox_ar, 0, wx.EXPAND)

        grid_conv.AddGrowableCol(1, 1)
        conv_sizer.Add(grid_conv, 1, wx.EXPAND | wx.ALL, 5)
        main_sizer.Add(conv_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # --- NEW: Boundary Conditions ---
        bc_box = wx.StaticBox(self, label=" Boundary Conditions (Domain Edges) ")
        bc_sizer = wx.StaticBoxSizer(bc_box, wx.VERTICAL)

        # PML Cells
        hbox_pml = wx.BoxSizer(wx.HORIZONTAL)
        hbox_pml.Add(wx.StaticText(bc_box, label="PML Absorber Cells (if PML chosen):"), 0,
                     wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.txt_pml_cells = wx.TextCtrl(bc_box, value="8")
        hbox_pml.Add(self.txt_pml_cells, 0, wx.EXPAND)
        bc_sizer.Add(hbox_pml, 0, wx.EXPAND | wx.BOTTOM, 10)

        grid_bc = wx.FlexGridSizer(3, 4, 8, 12)
        bc_choices = ["PML", "PEC", "PMC", "MUR"]

        # X-Axis
        grid_bc.Add(wx.StaticText(bc_box, label="-X (Left):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.cmb_bc_x_neg = wx.ComboBox(bc_box, choices=bc_choices, style=wx.CB_READONLY);
        self.cmb_bc_x_neg.SetValue("PML")
        grid_bc.Add(self.cmb_bc_x_neg, 0, wx.EXPAND)

        grid_bc.Add(wx.StaticText(bc_box, label="+X (Right):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.cmb_bc_x_pos = wx.ComboBox(bc_box, choices=bc_choices, style=wx.CB_READONLY);
        self.cmb_bc_x_pos.SetValue("PML")
        grid_bc.Add(self.cmb_bc_x_pos, 0, wx.EXPAND)

        # Y-Axis
        grid_bc.Add(wx.StaticText(bc_box, label="-Y (Bottom):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.cmb_bc_y_neg = wx.ComboBox(bc_box, choices=bc_choices, style=wx.CB_READONLY);
        self.cmb_bc_y_neg.SetValue("PML")
        grid_bc.Add(self.cmb_bc_y_neg, 0, wx.EXPAND)

        grid_bc.Add(wx.StaticText(bc_box, label="+Y (Top):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.cmb_bc_y_pos = wx.ComboBox(bc_box, choices=bc_choices, style=wx.CB_READONLY);
        self.cmb_bc_y_pos.SetValue("PML")
        grid_bc.Add(self.cmb_bc_y_pos, 0, wx.EXPAND)

        # Z-Axis
        grid_bc.Add(wx.StaticText(bc_box, label="-Z (Below Board):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.cmb_bc_z_neg = wx.ComboBox(bc_box, choices=bc_choices, style=wx.CB_READONLY);
        self.cmb_bc_z_neg.SetValue("PML")
        grid_bc.Add(self.cmb_bc_z_neg, 0, wx.EXPAND)

        grid_bc.Add(wx.StaticText(bc_box, label="+Z (Above Board):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.cmb_bc_z_pos = wx.ComboBox(bc_box, choices=bc_choices, style=wx.CB_READONLY);
        self.cmb_bc_z_pos.SetValue("PML")
        grid_bc.Add(self.cmb_bc_z_pos, 0, wx.EXPAND)

        bc_sizer.Add(grid_bc, 1, wx.EXPAND | wx.ALL, 5)
        main_sizer.Add(bc_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # --- Ambient Environment Settings ---
        ambient_box = wx.StaticBox(self, label=" Ambient Environment Settings ")
        ambient_sizer = wx.StaticBoxSizer(ambient_box, wx.VERTICAL)
        grid_ambient = wx.FlexGridSizer(2, 2, 8, 12)

        grid_ambient.Add(wx.StaticText(ambient_box, label="Ambient Epsilon (Er):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_ambient_epsilon = wx.TextCtrl(ambient_box, value="1.0")
        grid_ambient.Add(self.txt_ambient_epsilon, 0, wx.EXPAND)

        grid_ambient.Add(wx.StaticText(ambient_box, label="Ambient Mue (Ur):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_ambient_mue = wx.TextCtrl(ambient_box, value="1.0")
        grid_ambient.Add(self.txt_ambient_mue, 0, wx.EXPAND)

        grid_ambient.AddGrowableCol(1, 1)
        ambient_sizer.Add(grid_ambient, 1, wx.EXPAND | wx.ALL, 8)
        main_sizer.Add(ambient_sizer, 0, wx.EXPAND | wx.ALL, 10)

        # --- Advanced N-Port Matrix Execution ---
        adv_box = wx.StaticBox(self, label=" Advanced N-Port Matrix Execution ")
        adv_sizer = wx.StaticBoxSizer(adv_box, wx.VERTICAL)

        self.chk_is_pec = wx.CheckBox(adv_box, label="Use Ideal PEC for Copper (Lossless)")
        self.chk_is_pec.SetValue(False)
        self.chk_is_pec.SetToolTip(
            "When checked, copper layers use AddMetal('PEC'). Unchecked uses finite conductivity.")
        adv_sizer.Add(self.chk_is_pec, 0, wx.EXPAND | wx.ALL, 5)

        try:
            max_threads = os.cpu_count() or 4
        except Exception:
            max_threads = 4
        thread_choices = ["Auto"] + [str(i) for i in range(1, max_threads + 1)]

        hbox_threads = wx.BoxSizer(wx.HORIZONTAL)
        hbox_threads.Add(wx.StaticText(adv_box, label="FDTD Threads (Parallel):"), 0,
                         wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.choice_threads = wx.Choice(adv_box, choices=thread_choices)
        self.choice_threads.SetSelection(0)
        hbox_threads.Add(self.choice_threads, 0, wx.EXPAND)
        adv_sizer.Add(hbox_threads, 0, wx.EXPAND | wx.ALL, 5)

        # --- Group checkboxes horizontally to save vertical space ---
        hbox_checks = wx.BoxSizer(wx.HORIZONTAL)

        self.chk_parallel = wx.CheckBox(adv_box, label="Parallel Simulation (Multiprocessing)")
        self.chk_symmetry = wx.CheckBox(adv_box, label="Assume Reciprocal Symmetry")
        self.chk_symmetry.SetToolTip("S_ij=S_ji")
        self.chk_wsl2 = wx.CheckBox(adv_box, label="Use WSL2 Engine (Windows)")
        self.chk_wsl2.SetValue(True)

        hbox_checks.Add(self.chk_parallel, 1, wx.EXPAND | wx.RIGHT, 10)
        hbox_checks.Add(self.chk_symmetry, 1, wx.EXPAND | wx.LEFT, 10)
        hbox_checks.Add(self.chk_wsl2, 1, wx.EXPAND | wx.LEFT, 5)

        adv_sizer.Add(hbox_checks, 0, wx.EXPAND | wx.ALL, 8)
        main_sizer.Add(adv_sizer, 0, wx.EXPAND | wx.ALL, 10)

        self.SetSizer(main_sizer)
        self.Layout()

    def get_data(self):
        thread_val = self.choice_threads.GetStringSelection()
        num_threads = 0 if thread_val == "Auto" else int(thread_val)

        def safe_float(val, default=1.0):
            try:
                return float(val)
            except ValueError:
                return default

        return {
            "excitation_type": self.cmb_excite_type.GetValue(),
            "f_start_ghz": float(self.txt_f_start.GetValue()),
            "f_stop_ghz": float(self.txt_f_stop.GetValue()),
            "f_max_ghz": float(self.txt_f_max.GetValue()),
            "f_num_points": int(self.txt_num_points.GetValue()),
            "energy_limit_db": float(self.cmb_energy_limit.GetValue()),
            "max_timesteps": int(self.txt_max_timesteps.GetValue()),
            "use_ar_filter": self.chk_ar_filter.IsChecked(),
            "ar_extrap_steps": int(self.txt_ar_steps.GetValue()) if self.txt_ar_steps.GetValue().isdigit() else 100000,
            "pml_cells": int(self.txt_pml_cells.GetValue()),
            "boundary_conditions": {
                "x_neg": self.cmb_bc_x_neg.GetValue(),
                "x_pos": self.cmb_bc_x_pos.GetValue(),
                "y_neg": self.cmb_bc_y_neg.GetValue(),
                "y_pos": self.cmb_bc_y_pos.GetValue(),
                "z_neg": self.cmb_bc_z_neg.GetValue(),
                "z_pos": self.cmb_bc_z_pos.GetValue()
            },
            "ambient_epsilon": safe_float(self.txt_ambient_epsilon.GetValue(), 1.0),
            "ambient_mue": safe_float(self.txt_ambient_mue.GetValue(), 1.0),
            "use_pec_copper": self.chk_is_pec.IsChecked(),
            "threads": num_threads,
            "assume_symmetry": self.chk_symmetry.GetValue(),
            "parallel_sim": self.chk_parallel.GetValue(),
            "use_wsl2": self.chk_wsl2.GetValue()
        }

class MeshSettingsTab(wx.Panel):
    def __init__(self, parent, main_app_ref=None):
        super().__init__(parent)
        self.main_app_ref = main_app_ref
        main_sizer = wx.BoxSizer(wx.VERTICAL)

        '''
        self.chk_enable_crop = wx.CheckBox(domain_box, label="Enable Trace Cropping (X/Y)")
        self.chk_enable_crop.SetValue(True)
        domain_grid.Add(self.chk_enable_crop, 0, wx.ALIGN_CENTER_VERTICAL)
        '''

        # --- 0. Domain & Air Margins ---
        domain_box = wx.StaticBox(self, label=" 0. Domain Boundaries & Air Margins (mm) ")
        domain_sizer = wx.StaticBoxSizer(domain_box, wx.VERTICAL)

        # 1. Add the checkbox to the vertical sizer so it sits neatly at the top
        self.chk_enable_crop = wx.CheckBox(domain_box, label="Enable Trace Cropping (X/Y)")
        self.chk_enable_crop.SetValue(True)
        domain_sizer.Add(self.chk_enable_crop, 0, wx.BOTTOM | wx.LEFT, 8)

        # 2. Set rows to 0 (dynamic) and cols to 4 to prevent future capacity limits
        domain_grid = wx.FlexGridSizer(0, 4, 8, 12)

        # X-Axis
        domain_grid.Add(wx.StaticText(domain_box, label="-X Margin (Left):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_margin_x_neg = wx.TextCtrl(domain_box, value="5.0")
        domain_grid.Add(self.txt_margin_x_neg, 0, wx.EXPAND)

        domain_grid.Add(wx.StaticText(domain_box, label="+X Margin (Right):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_margin_x_pos = wx.TextCtrl(domain_box, value="5.0")
        domain_grid.Add(self.txt_margin_x_pos, 0, wx.EXPAND)

        # Y-Axis
        domain_grid.Add(wx.StaticText(domain_box, label="-Y Margin (Bottom):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_margin_y_neg = wx.TextCtrl(domain_box, value="5.0")
        domain_grid.Add(self.txt_margin_y_neg, 0, wx.EXPAND)

        domain_grid.Add(wx.StaticText(domain_box, label="+Y Margin (Top):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_margin_y_pos = wx.TextCtrl(domain_box, value="5.0")
        domain_grid.Add(self.txt_margin_y_pos, 0, wx.EXPAND)

        # Z-Axis
        domain_grid.Add(wx.StaticText(domain_box, label="-Z Margin (Below):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_margin_z_neg = wx.TextCtrl(domain_box, value="4.0")
        domain_grid.Add(self.txt_margin_z_neg, 0, wx.EXPAND)

        domain_grid.Add(wx.StaticText(domain_box, label="+Z Margin (Above):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_margin_z_pos = wx.TextCtrl(domain_box, value="4.0")
        domain_grid.Add(self.txt_margin_z_pos, 0, wx.EXPAND)

        domain_grid.AddGrowableCol(1, 1)
        domain_grid.AddGrowableCol(3, 1)

        # Add the 12-item grid to the main domain sizer below the checkbox
        domain_sizer.Add(domain_grid, 1, wx.EXPAND | wx.ALL, 5)
        main_sizer.Add(domain_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # --- 1. Mode Selection ---
        mode_box = wx.BoxSizer(wx.HORIZONTAL)
        mode_box.Add(wx.StaticText(self, label="Mesh Mode:"), 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.cmb_mode = wx.ComboBox(self, choices=["Automatic", "Advanced"], style=wx.CB_READONLY)
        self.cmb_mode.SetValue("Advanced")
        mode_box.Add(self.cmb_mode, 0, wx.EXPAND)
        main_sizer.Add(mode_box, 0, wx.EXPAND | wx.ALL, 5)

        # --- 2. Global / Background Resolution ---
        global_box = wx.StaticBox(self, label=" 1. Global / Background Resolution ")
        global_sizer = wx.StaticBoxSizer(global_box, wx.VERTICAL)
        grid_global = wx.FlexGridSizer(4, 4, 8, 12)

        # Headers (Parent is now global_box to fix wxPython warnings)
        grid_global.Add(wx.StaticText(global_box, label="Parameter"), 0, wx.ALIGN_CENTER_VERTICAL)
        grid_global.Add(wx.StaticText(global_box, label="X Axis"), 0, wx.ALIGN_CENTER)
        grid_global.Add(wx.StaticText(global_box, label="Y Axis"), 0, wx.ALIGN_CENTER)
        grid_global.Add(wx.StaticText(global_box, label="Z Axis"), 0, wx.ALIGN_CENTER)

        # Max Cell Size
        grid_global.Add(wx.StaticText(global_box, label="Max cell size (mm):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_global_x = wx.TextCtrl(global_box, value="1.20")
        self.txt_global_y = wx.TextCtrl(global_box, value="1.20")
        self.txt_global_z = wx.TextCtrl(global_box, value="0.50")
        grid_global.Add(self.txt_global_x, 0, wx.EXPAND)
        grid_global.Add(self.txt_global_y, 0, wx.EXPAND)
        grid_global.Add(self.txt_global_z, 0, wx.EXPAND)

        # Growth Characteristic
        grid_global.Add(wx.StaticText(global_box, label="Growth Type:"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.cmb_growth_x = wx.ComboBox(global_box, choices=["Linear", "Exponential"], style=wx.CB_READONLY,
                                        value="Exponential")
        self.cmb_growth_y = wx.ComboBox(global_box, choices=["Linear", "Exponential"], style=wx.CB_READONLY,
                                        value="Exponential")
        self.cmb_growth_z = wx.ComboBox(global_box, choices=["Linear", "Exponential"], style=wx.CB_READONLY,
                                        value="Exponential")
        grid_global.Add(self.cmb_growth_x, 0, wx.EXPAND)
        grid_global.Add(self.cmb_growth_y, 0, wx.EXPAND)
        grid_global.Add(self.cmb_growth_z, 0, wx.EXPAND)

        # Max Cell Ratio
        grid_global.Add(wx.StaticText(global_box, label="Max cell ratio:"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_ratio_x = wx.TextCtrl(global_box, value="1.30")
        self.txt_ratio_y = wx.TextCtrl(global_box, value="1.30")
        self.txt_ratio_z = wx.TextCtrl(global_box, value="1.20")
        grid_global.Add(self.txt_ratio_x, 0, wx.EXPAND)
        grid_global.Add(self.txt_ratio_y, 0, wx.EXPAND)
        grid_global.Add(self.txt_ratio_z, 0, wx.EXPAND)

        global_sizer.Add(grid_global, 1, wx.EXPAND | wx.ALL, 5)
        main_sizer.Add(global_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # --- 3. Feature & Gap Resolution ---
        feat_box = wx.StaticBox(self, label=" 2 & 3. Feature & Gap Resolution ")
        feat_sizer = wx.StaticBoxSizer(feat_box, wx.VERTICAL)
        grid_feat = wx.FlexGridSizer(4, 4, 8, 12)

        grid_feat.Add(wx.StaticText(feat_box, label=""), 0, wx.ALIGN_CENTER_VERTICAL)
        grid_feat.Add(wx.StaticText(feat_box, label="X Axis"), 0, wx.ALIGN_CENTER)
        grid_feat.Add(wx.StaticText(feat_box, label="Y Axis"), 0, wx.ALIGN_CENTER)
        grid_feat.Add(wx.StaticText(feat_box, label="Z Axis"), 0, wx.ALIGN_CENTER)

        grid_feat.Add(wx.StaticText(feat_box, label="Min cells across feature:"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_feat_x = wx.TextCtrl(feat_box, value="3")
        self.txt_feat_y = wx.TextCtrl(feat_box, value="3")
        self.txt_feat_z = wx.TextCtrl(feat_box, value="2")
        grid_feat.Add(self.txt_feat_x, 0, wx.EXPAND)
        grid_feat.Add(self.txt_feat_y, 0, wx.EXPAND)
        grid_feat.Add(self.txt_feat_z, 0, wx.EXPAND)

        grid_feat.Add(wx.StaticText(feat_box, label="Min cells in gap:"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_gap_x = wx.TextCtrl(feat_box, value="3")
        self.txt_gap_y = wx.TextCtrl(feat_box, value="3")
        self.txt_gap_z = wx.TextCtrl(feat_box, value="1")
        grid_feat.Add(self.txt_gap_x, 0, wx.EXPAND)
        grid_feat.Add(self.txt_gap_y, 0, wx.EXPAND)
        grid_feat.Add(self.txt_gap_z, 0, wx.EXPAND)

        # Min cell size safety
        grid_feat.Add(wx.StaticText(feat_box, label="Min cell size limit (mm):"), 0, wx.ALIGN_CENTER_VERTICAL)
        self.txt_min_cell = wx.TextCtrl(feat_box, value="0.04")
        grid_feat.Add(self.txt_min_cell, 0, wx.EXPAND)
        grid_feat.Add(wx.StaticText(feat_box, label=""), 0, wx.EXPAND)
        grid_feat.Add(wx.StaticText(feat_box, label=""), 0, wx.EXPAND)

        feat_sizer.Add(grid_feat, 1, wx.EXPAND | wx.ALL, 5)
        main_sizer.Add(feat_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # --- 4. Diagonal Trace Resolution ---
        diag_box = wx.StaticBox(self, label=" 4. Diagonal Trace Resolution ")
        diag_sizer = wx.StaticBoxSizer(diag_box, wx.VERTICAL)

        self.chk_mesh_diagonals = wx.CheckBox(diag_box, label="Force Dense Mesh on Diagonal Traces")
        self.chk_mesh_diagonals.SetValue(True)
        self.chk_mesh_diagonals.SetToolTip("Uncheck to rely entirely on the Global Background Mesh (faster).")
        diag_sizer.Add(self.chk_mesh_diagonals, 0, wx.ALL, 5)

        hbox_diag = wx.BoxSizer(wx.HORIZONTAL)
        hbox_diag.Add(wx.StaticText(diag_box, label="Diagonal Cells per Trace Width:"), 0,
                      wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 5)
        self.txt_trace_cells = wx.TextCtrl(diag_box, value="1")
        self.txt_trace_cells.SetToolTip("1 is usually enough to connect the staircase. 3+ will cause overmeshing.")
        hbox_diag.Add(self.txt_trace_cells, 0, wx.EXPAND)
        diag_sizer.Add(hbox_diag, 0, wx.EXPAND | wx.ALL, 5)

        main_sizer.Add(diag_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # --- 5. Geometry Preservation ---
        geom_box = wx.StaticBox(self, label=" 5. Geometry Preservation (Hard Constraints) ")
        geom_sizer = wx.StaticBoxSizer(geom_box, wx.VERTICAL)

        self.chk_cond = wx.CheckBox(geom_box, label="Lock Conductor boundaries")
        self.chk_cond.SetValue(True)
        self.chk_via = wx.CheckBox(geom_box, label="Lock Via boundaries")
        self.chk_via.SetValue(True)
        self.chk_pad = wx.CheckBox(geom_box, label="Lock Pad boundaries")
        self.chk_pad.SetValue(True)
        self.chk_diel = wx.CheckBox(geom_box, label="Lock Dielectric boundaries")
        self.chk_diel.SetValue(True)
        self.chk_port = wx.CheckBox(geom_box, label="Lock Port boundaries")
        self.chk_port.SetValue(True)

        geom_grid = wx.GridSizer(2, 3, 5, 5)
        geom_grid.AddMany([self.chk_cond, self.chk_via, self.chk_pad, self.chk_diel, self.chk_port])
        geom_sizer.Add(geom_grid, 0, wx.EXPAND | wx.ALL, 5)
        main_sizer.Add(geom_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # --- 5. Action Buttons ---
        self.btn_calc_mesh = wx.Button(self, label="🔍 Calculate Mesh & Show Cell Count")
        self.btn_calc_mesh.SetBackgroundColour(wx.Colour(230, 240, 255))
        self.btn_calc_mesh.Bind(wx.EVT_BUTTON, self.on_calculate_mesh)
        main_sizer.Add(self.btn_calc_mesh, 0, wx.EXPAND | wx.ALL, 10)

        self.SetSizer(main_sizer)
        self.Layout()

    def on_calculate_mesh(self, event):
        main_dialog = getattr(self, 'main_app_ref', None)
        if not main_dialog:
            main_dialog = self.GetParent().GetParent()

        # 1. Export the latest data to JSON before running the script
        if main_dialog and hasattr(main_dialog, '_export_json'):
            main_dialog._export_json()

        json_path = getattr(main_dialog, 'last_export_path', None)
        if not json_path or not os.path.exists(json_path):
            wx.MessageBox("Please 'Export to JSON' first.", "Error", wx.OK | wx.ICON_ERROR)
            return

        script_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "solvers", "run_openems.py")
        cmd = [sys.executable, script_path, json_path, "--calc-mesh"]

        btn = event.GetEventObject()
        btn.Disable()
        main_dialog.txt_output.AppendText("\n[*] Calculating FDTD Mesh bounds and estimating cell count...\n")
        wx.Yield()

        try:
            startupinfo = None
            if os.name == 'nt':
                startupinfo = subprocess.STARTUPINFO()
                startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
                creationflags = subprocess.CREATE_NO_WINDOW

            result = subprocess.run(cmd, capture_output=True, text=True, startupinfo=startupinfo, check=True, creationflags=creationflags)
            output = result.stdout
            main_dialog.txt_output.AppendText(output)

            # --- NEW REGEX PARSERS FOR THE UPDATED DIAGNOSTICS ---
            total_match = re.search(r'\[\*\] Total FDTD Cells:\s*([\d,]+)\s*\((.*?)\)', output)
            min_cell_match = re.search(r'\[\*\] Absolute Smallest Cell:\s*([0-9.]+)\s*mm', output)

            if total_match:
                total_cells = total_match.group(1)
                dimensions = total_match.group(2)
                min_cell = min_cell_match.group(1) if min_cell_match else "Unknown"
                self.mesh_cells_str = total_cells

                # Format the message box popup
                msg = f"Mesh calculation successful.\n\nTotal FDTD Cells: {total_cells}\nGrid Dimensions: {dimensions}\nSmallest Cell: {min_cell} mm"

                # Check if our Python engine threw any warnings
                if "[FAILED]" in output:
                    msg += "\n\n[!] Warnings detected. Check the log for details."
                    wx.MessageBox(msg, "Mesh Estimation (Warnings)", wx.OK | wx.ICON_WARNING, self)
                else:
                    wx.MessageBox(msg, "Mesh Estimation", wx.OK | wx.ICON_INFORMATION, self)
            else:
                wx.MessageBox("Could not parse mesh summary from output. Check the log.", "Error",
                              wx.OK | wx.ICON_ERROR, self)

        except subprocess.CalledProcessError as e:
            wx.MessageBox(f"Failed to calculate mesh.\nError code: {e.returncode}", "Error", wx.OK | wx.ICON_ERROR,
                          self)
        finally:
            btn.Enable()

    def get_data(self):
        def safe_float(val, default=0.0):
            try:
                return float(val)
            except ValueError:
                return default
        return {
            "mesh_mode": self.cmb_mode.GetValue(),
            "crop_to_active_traces": self.chk_enable_crop.IsChecked(),
            "air_margins_mm": {
                "x_neg": safe_float(self.txt_margin_x_neg.GetValue(), 5.0),
                "x_pos": safe_float(self.txt_margin_x_pos.GetValue(), 5.0),
                "y_neg": safe_float(self.txt_margin_y_neg.GetValue(), 5.0),
                "y_pos": safe_float(self.txt_margin_y_pos.GetValue(), 5.0),
                "z_neg": safe_float(self.txt_margin_z_neg.GetValue(), 4.0),
                "z_pos": safe_float(self.txt_margin_z_pos.GetValue(), 4.0)
            },
            "min_cell_size_mm": float(self.txt_min_cell.GetValue()),
            "mesh_diagonal_traces": self.chk_mesh_diagonals.IsChecked(),
            "mesh_global": {
                "x": {"max_size": float(self.txt_global_x.GetValue()), "growth": self.cmb_growth_x.GetValue(),
                      "ratio": float(self.txt_ratio_x.GetValue())},
                "y": {"max_size": float(self.txt_global_y.GetValue()), "growth": self.cmb_growth_y.GetValue(),
                      "ratio": float(self.txt_ratio_y.GetValue())},
                "z": {"max_size": float(self.txt_global_z.GetValue()), "growth": self.cmb_growth_z.GetValue(),
                      "ratio": float(self.txt_ratio_z.GetValue())},
            },
            "mesh_feature": {
                "x_cells": int(self.txt_feat_x.GetValue()),
                "y_cells": int(self.txt_feat_y.GetValue()),
                "z_cells": int(self.txt_feat_z.GetValue()),
                "trace_cells": int(self.txt_trace_cells.GetValue()) if self.txt_trace_cells.GetValue().strip() else 1
            },
            "mesh_gap": {
                "x_cells": int(self.txt_gap_x.GetValue()),
                "y_cells": int(self.txt_gap_y.GetValue()),
                "z_cells": int(self.txt_gap_z.GetValue())
            },
            "mesh_locks": {
                "conductor": self.chk_cond.IsChecked(),
                "via": self.chk_via.IsChecked(),
                "pad": self.chk_pad.IsChecked(),
                "dielectric": self.chk_diel.IsChecked(),
                "port": self.chk_port.IsChecked()
            }
        }

class DynamicPortsTab(wx.Panel):
    def __init__(self, parent, main_app):
        super().__init__(parent)
        self.main_app = main_app
        self.configured_ports = []

        main_sizer = wx.BoxSizer(wx.VERTICAL)

        # Header area
        header_sizer = wx.BoxSizer(wx.HORIZONTAL)
        self.lbl_count = wx.StaticText(self, label="Configured Ports: 0")
        self.lbl_count.SetFont(wx.Font(11, wx.FONTFAMILY_DEFAULT, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_BOLD))
        header_sizer.Add(self.lbl_count, 1, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 5)

        btn_add = wx.Button(self, label="➕ Add Port")
        btn_add.SetBackgroundColour(wx.Colour(200, 255, 200))
        btn_add.Bind(wx.EVT_BUTTON, self.on_add_port)
        header_sizer.Add(btn_add, 0, wx.ALL, 5)
        main_sizer.Add(header_sizer, 0, wx.EXPAND | wx.ALL, 5)

        # Scrollable list area
        self.scroll_win = wx.ScrolledWindow(self, style=wx.VSCROLL)
        self.scroll_win.SetScrollRate(5, 5)
        self.list_sizer = wx.BoxSizer(wx.VERTICAL)
        self.scroll_win.SetSizer(self.list_sizer)

        main_sizer.Add(self.scroll_win, 1, wx.EXPAND | wx.ALL, 5)
        self.SetSizer(main_sizer)

    def refresh_list(self):
        self.list_sizer.Clear(True)
        self.lbl_count.SetLabel(f"Configured Ports: {len(self.configured_ports)}")

        for idx, port_data in enumerate(self.configured_ports):
            row_panel = wx.Panel(self.scroll_win)
            row_panel.SetBackgroundColour(wx.Colour(240, 240, 240))
            row_sizer = wx.BoxSizer(wx.HORIZONTAL)

            mode = port_data.get('mode', '')
            net_p = port_data.get('positive_terminal', {}).get('net', 'None')
            info_str = f"Port {idx + 1} | {port_data.get('type', '')} | {mode}\n(+) {net_p}"
            if mode != "Single-Ended":
                net_n = port_data.get('negative_terminal', {}).get('net', 'None')
                info_str += f"  (-) {net_n}"

            lbl = wx.StaticText(row_panel, label=info_str)
            row_sizer.Add(lbl, 1, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 5)

            btn_edit = wx.Button(row_panel, label="Edit")
            btn_edit.Bind(wx.EVT_BUTTON, lambda evt, i=idx: self.on_edit_port(i))
            row_sizer.Add(btn_edit, 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)

            btn_rem = wx.Button(row_panel, label="Remove")
            btn_rem.SetForegroundColour(wx.RED)
            btn_rem.Bind(wx.EVT_BUTTON, lambda evt, i=idx: self.on_remove_port(i))
            row_sizer.Add(btn_rem, 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)

            row_panel.SetSizer(row_sizer)
            self.list_sizer.Add(row_panel, 0, wx.EXPAND | wx.BOTTOM, 5)

        self.scroll_win.Layout()
        self.scroll_win.Refresh()

    def on_add_port(self, event):
        copper_layers = getattr(self.main_app, 'active_copper_layers', ["F.Cu", "B.Cu"])
        dlg = PortConfigDialog(self, self.main_app.get_net_names(), self.main_app.get_pads_for_net, copper_layers)
        if dlg.ShowModal() == wx.ID_OK:
            self.configured_ports.append(dlg.get_port_data())
            self.refresh_list()
        dlg.Destroy()

    def on_edit_port(self, idx):
        copper_layers = getattr(self.main_app, 'active_copper_layers', ["F.Cu", "B.Cu"])
        dlg = PortConfigDialog(self, self.main_app.get_net_names(), self.main_app.get_pads_for_net, copper_layers, self.configured_ports[idx])
        if dlg.ShowModal() == wx.ID_OK:
            self.configured_ports[idx] = dlg.get_port_data()
            self.refresh_list()
        dlg.Destroy()

    def on_remove_port(self, idx):
        self.configured_ports.pop(idx)
        self.refresh_list()


class TelegramSimBot:
    def __init__(self, bot_token, chat_id, status_file_path):
        self.bot_token = bot_token
        self.chat_id = str(chat_id)
        self.status_file_path = status_file_path
        self.base_url = f"https://api.telegram.org/bot{self.bot_token}"
        self.is_listening = False
        self.last_update_id = 0

    def start_listening(self):
        """Spawns a daemon thread to listen for user commands without blocking the GUI."""
        self.is_listening = bool(len(self.bot_token))
        try:
            resp = requests.get(f"{self.base_url}/getMe", timeout=10)
            if resp.ok:
                print(f"[Telegram] Token configured and valid: YES")
                print(f"[Telegram] Chat ID configured: {self.chat_id}")
                print("[Telegram] Starting polling...")
            else:
                print(f"[Telegram] Token validation failed: {resp.status_code} {resp.text}")
                self.is_listening = False
                return
        except Exception as e:
            print(f"[Telegram] Failed to connect during initialization: {e}")
            self.is_listening = False
            return

        thread = threading.Thread(target=self._poll_updates, daemon=True)
        thread.start()

    def _poll_updates(self):
        """Robust long-polling loop with strict exception and timeout handling."""
        while self.is_listening:
            try:
                url = f"{self.base_url}/getUpdates?offset={self.last_update_id + 1}&timeout=30"
                response = requests.get(url, timeout=35)

                if not response.ok:
                    print(f"[Telegram] getUpdates failed: {response.status_code} {response.text}")
                    time.sleep(3)
                    continue

                data = response.json()
                if data.get("ok"):
                    for result in data.get("result", []):
                        self.last_update_id = result["update_id"]
                        message = result.get("message", {})
                        text = message.get("text", "")
                        chat_id = str(message.get("chat", {}).get("id"))

                        if text in ["/status", "\\status"]:
                            if chat_id == self.chat_id:
                                self._reply_with_status()
                            else:
                                print(
                                    f"[Telegram] Ignored /status from unauthorized chat_id={chat_id} (Configured={self.chat_id})")

            except requests.exceptions.RequestException as e:
                print(f"[Telegram] Network error during polling: {e}")
                time.sleep(3)
            except Exception as e:
                print(f"[Telegram] Unexpected error: {e}")
                time.sleep(3)

            time.sleep(1)

    def _reply_with_status(self):
        """Reads the live GUI dashboard JSON and safely sends it back."""
        if not self.is_listening:
            return
        try:
            if os.path.exists(self.status_file_path):
                with open(self.status_file_path, "r", encoding="utf-8") as f:
                    data = json.load(f)

                status = data.get("text_status", "Running")
                pct = data.get("progress_pct", 0)
                eta = data.get("eta", "ETA: Calculating...")
                energy = data.get("energy_db", 0.0)

                msg = f"📊 *Live Simulation Status*\n\n*Status:* {status}\n*Progress:* {pct}%\n*Energy:* {energy:.2f} dB\n*Time:* {eta}"
            else:
                msg = "⚠️ Simulation JSON not found. Engine might be initializing."

            resp = requests.post(
                f"{self.base_url}/sendMessage",
                json={"chat_id": self.chat_id, "text": msg, "parse_mode": "Markdown"},
                timeout=15
            )
            if not resp.ok:
                print(f"[Telegram] sendMessage failed: {resp.status_code} {resp.text}")

        except requests.exceptions.RequestException as e:
            print(f"[Telegram] Failed to send status reply: {e}")
        except Exception as e:
            print(f"⚠️ Error parsing live status: {e}")

    def send_completion(self, plot_path=None):
        """Sends the final notification and image via a NON-DAEMON background thread."""
        if not self.is_listening:
            return

        def _upload():
            msg = "✅ *Simulation Complete!* Your S-parameters are ready."
            try:
                if plot_path and os.path.exists(plot_path):
                    with open(plot_path, "rb") as photo:
                        response = requests.post(
                            f"{self.base_url}/sendPhoto",
                            data={"chat_id": self.chat_id, "caption": msg, "parse_mode": "Markdown"},
                            files={"photo": photo},
                            timeout=30
                        )
                else:
                    response = requests.post(
                        f"{self.base_url}/sendMessage",
                        json={"chat_id": self.chat_id, "text": msg, "parse_mode": "Markdown"},
                        timeout=15
                    )
                if not response.ok:
                    print(f"[Telegram] Completion upload failed: {response.status_code} {response.text}")
            except Exception as e:
                print(f"[!] Failed to upload Telegram plot: {e}")

        # daemon=False ensures KiCad won't immediately terminate the upload if closed
        threading.Thread(target=_upload, daemon=False).start()

class SimulationThread(threading.Thread):
    def __init__(self, cmd, log_callback, done_callback):
        super().__init__()
        self.cmd = cmd
        self.log_callback = log_callback
        self.done_callback = done_callback
        self.process = None
        self.abort_flag = False

    def run(self):
        try:
            startupinfo = None
            creationflags = 0
            if os.name == 'nt':
                startupinfo = subprocess.STARTUPINFO()
                startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
                creationflags = subprocess.CREATE_NO_WINDOW

            self.process = subprocess.Popen(
                self.cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                universal_newlines=True,
                encoding='utf-8',
                errors='replace',
                startupinfo=startupinfo,
                creationflags=creationflags
            )

            buffer = []
            last_flush = time.time()

            while True:
                if self.abort_flag:
                    break
                line = self.process.stdout.readline()

                if not line and self.process.poll() is not None:
                    if buffer:
                        wx.CallAfter(self.log_callback, "".join(buffer))
                    break

                if line:
                    buffer.append(line)
                    # Flush buffer to GUI every 0.1s or 100 lines to prevent Event Loop freezing
                    if time.time() - last_flush > 0.1 or len(buffer) > 100:
                        wx.CallAfter(self.log_callback, "".join(buffer))
                        buffer = []
                        last_flush = time.time()

            if self.process is not None:
                self.process.wait()


        except Exception as e:
            err_msg = traceback.format_exc()
            wx.CallAfter(self.log_callback, f"Process Error:\n{err_msg}\n")
        finally:
            success = not self.abort_flag and (self.process and self.process.returncode == 0)
            wx.CallAfter(self.done_callback, success)

    def abort(self):
        # MANUAL ABORT: Hard kill, do not save results.
        self.abort_flag = True
        if self.process:
            try:
                if os.name == 'nt':
                    subprocess.run(["taskkill", "/T", "/F", "/PID", str(self.process.pid)],
                                   capture_output=True, creationflags=subprocess.CREATE_NO_WINDOW)
                else:
                    self.process.kill()
            except Exception:
                self.process.terminate()
            self.process = None

    def graceful_stop(self, sim_dir):
        """
        Creates an 'ABORT' file. The openEMS C++ core natively polls for this file,
        and when found, cleanly breaks the time-stepping loop, saves fields, and continues.
        """
        if not sim_dir: return
        abort_path = os.path.join(sim_dir, "ABORT")
        try:
            with open(abort_path, 'w') as f:
                f.write("STOP")
        except Exception as e:
            print(f"Graceful stop failed to create ABORT file: {e}")

if __name__ == '__main__':
    if os.name == 'nt':
        try:
            myappid = 'com.sim.em_studio.gui.1'  # Arbitrary unique identifier
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(myappid)
        except Exception:
            pass
        openems_locator()

    app = wx.App()
    try:
        # kipy automatically reads KICAD_API_SOCKET from the environment when launched by KiCad
        socket_path = os.environ.get("KICAD_API_SOCKET")

        # Initialize the client (adding a timeout prevents infinite hanging if KiCad is busy)
        kicad = KiCad(socket_path=socket_path, timeout_ms=3000)
        board = kicad.get_board()

        dlg = EMSimDialog(None, kicad, board)
        dlg.ShowModal()
        dlg.Destroy()

    except Exception as e:
        wx.MessageBox(
            f"Failed to connect to KiCad IPC API.\nEnsure PCB Editor is open and API is enabled.\n\nError: {e}",
            "Connection Error", wx.OK | wx.ICON_ERROR
        )
